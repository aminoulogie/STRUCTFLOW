"""Render a Trace as a Word note de calcul (python-docx)."""
from __future__ import annotations

import re
from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .trace import Figure, Para, Step, Table, Trace

GREEN = RGBColor(0x00, 0x80, 0x3C)
RED = RGBColor(0xC0, 0x00, 0x00)
GREY = RGBColor(0x60, 0x60, 0x60)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
MATH = "Cambria Math"


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _run(p, text, bold=False, italic=False, size=None, color=None, font=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def _page_number(paragraph):
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = paragraph.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)


EQUATIONS = {"on": True, "num": True}
# columns printed for the current chapter: formula, numerical application, result, verification
COLS = {"sym": True, "num": True, "res": True, "ok": True}


def _eq(q, text):
    """Append text as a native Word equation (LaTeX → OMML). Returns False if it should stay plain text."""
    if not EQUATIONS["on"]:
        return False
    try:
        from . import mathfmt
        if not mathfmt.available():
            return False
        latex = mathfmt.to_latex(text)
        words = re.findall(r"\\text\{([^}]*)\}", latex)
        if sum(1 for w in words if len(w) >= 4) > 2:  # mostly French text: keep it readable as plain text
            return False
        om = mathfmt.latex_to_omml(latex)
        if q.runs:
            q._p.append(om)  # inline after text
        else:  # display equation (full-size fractions), left-aligned
            para = OxmlElement("m:oMathPara")
            pr = OxmlElement("m:oMathParaPr")
            jc = OxmlElement("m:jc")
            jc.set(qn("m:val"), "left")
            pr.append(jc)
            para.append(pr)
            para.append(om)
            q._p.append(para)
        return True
    except Exception:  # noqa: BLE001
        return False


def _shown(s: Step):
    """False when the chapter's ticked columns leave nothing to print for this step."""
    return bool((COLS["sym"] and s.sym) or (COLS["num"] and s.num) or (COLS["res"] and s.res)
                or (COLS["ok"] and s.ok is not None))


def _step(doc, s: Step):
    if not _shown(s):
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    _run(p, "▪ " + s.label, bold=True, size=10)
    if s.ref:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(0.6)
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.keep_with_next = True
        _run(q, "Réf. : " + s.ref, italic=True, size=8, color=GREY)
    lines = []
    shown_sym = bool(s.sym and COLS["sym"])
    if shown_sym:
        lines.append((s.sym, False))

    def pre(text, cond):
        return ("= " if cond and "=" not in text and "→" not in text[:3] else "") + text
    if s.num and COLS["num"]:
        lines.append((pre(s.num, shown_sym), False))
    if s.res and COLS["res"]:
        lines.append((pre(s.res, shown_sym or bool(s.num and COLS["num"])), True))
    for text, bold in lines:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(1.0)
        q.paragraph_format.space_after = Pt(1)
        if not _eq(q, ("⇒ " + (text[2:] if text.startswith("= ") else text)) if bold else text):
            if bold:
                _run(q, "⇒ ", bold=True, size=10, color=BLUE)
            _run(q, text, bold=bold, size=10, font=MATH, color=BLUE if bold else None)
    if s.ok is not None and COLS["ok"]:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(1.0)
        q.paragraph_format.space_after = Pt(2)
        if s.cond:
            _run(q, "Condition : ", size=10)
            if not _eq(q, s.cond):
                _run(q, s.cond, size=10, font=MATH)
            _run(q, "  →  ", size=10)
        _run(q, "VÉRIFIÉE ✔" if s.ok else "NON VÉRIFIÉE ✘", bold=True, size=10, color=GREEN if s.ok else RED)


def _table(doc, t: Table):
    if t.title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        _run(p, t.title, bold=True, italic=True, size=9)
    ncol = len(t.headers)
    tab = doc.add_table(rows=1 if t.headers and any(t.headers) else 0, cols=ncol)
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    small = ncol > 6 or len(t.rows) > 30
    fs = 7 if small else 9
    if t.headers and any(t.headers):
        for i, h in enumerate(t.headers):
            c = tab.rows[0].cells[i]
            c.text = ""
            _run(c.paragraphs[0], str(h), bold=True, size=fs)
            _shade(c, "D9E2F3")
    for row in t.rows:
        cells = tab.add_row().cells
        if isinstance(row, dict) and "_group" in row:  # merged, shaded group header
            m = cells[0].merge(cells[-1])
            m.text = ""
            p = m.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, row["_group"], bold=True, size=fs + 1)
            _shade(m, "BFBFBF" if row.get("dark", True) else "D9D9D9")
            continue
        for i, v in enumerate(row):
            cells[i].text = ""
            txt = str(v)
            color = RED if txt.startswith("NON OK") else (GREEN if txt == "OK" else None)
            _run(cells[i].paragraphs[0], txt, size=fs, color=color, bold=color is not None)
    doc.add_paragraph()


# ---------------------------------------------------------------- "Tableau Excel" style
XL_FONT = "Times New Roman"
XL_HEAD = "BFBFBF"      # grey header cells
XL_LABEL = "F2F2F2"     # light-grey row labels
XL_RES = "00B050"       # green result cells
XL_RED = RGBColor(0xFF, 0x00, 0x00)


def _xl_borders(tab):
    """Double borders on every cell edge, like the user's calculation sheets."""
    tblPr = tab._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "double")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)


def _xl_cell(cell, text, fill=None, color=None, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    _run(p, str(text), bold=bold, italic=italic, size=size, color=color, font=XL_FONT)
    if fill:
        _shade(cell, fill)


def _xl_ok(txt):
    t = str(txt)
    if t.startswith("NON") or "✘" in t:
        return XL_RED, "FFC7CE"
    if t == "OK" or t.startswith("OK") or "✔" in t:
        return XL_RED, None
    return None, None


def _xl_table(doc, t: Table):
    """Any result table drawn like an Excel sheet."""
    if t.title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        _run(p, t.title, bold=True, size=10, font=XL_FONT)
    ncol = len(t.headers)
    tab = doc.add_table(rows=1 if t.headers and any(t.headers) else 0, cols=ncol)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    _xl_borders(tab)
    fs = 7 if ncol > 7 or len(t.rows) > 30 else 8.5
    if t.headers and any(t.headers):
        for i, h in enumerate(t.headers):
            _xl_cell(tab.rows[0].cells[i], h, XL_HEAD, size=fs)
    for row in t.rows:
        cells = tab.add_row().cells
        if isinstance(row, dict) and "_group" in row:
            m = cells[0].merge(cells[-1])
            _xl_cell(m, row["_group"], XL_HEAD if row.get("dark", True) else "D9D9D9", size=fs + 1)
            continue
        for i, v in enumerate(row):
            color, fill = _xl_ok(v)
            if i == 0 and ncol > 2:
                fill = fill or XL_LABEL
            _xl_cell(cells[i], v, fill, color, size=fs)
    doc.add_paragraph()


GRID_COLS = (("sym", "Formule", 4.4), ("num", "Application numérique", 4.0), ("res", "Résultat", 2.8),
             ("ok", "Vérification", 1.9))


def _grid(doc, steps, excel=False):
    """'Tableaux' and 'Tableau Excel' styles: consecutive steps as one table, with only the ticked columns."""
    steps = [st for st in steps if _shown(st)]
    if not steps:
        return
    cols = [c for c in GRID_COLS if COLS[c[0]]]
    used = sum(w for _, _, w in cols)
    lab_w = max(4.4, min(8.0, 17.0 - used))
    k = (17.0 - lab_w) / used if used else 1.0  # spread the remaining width over the ticked columns
    heads = [("label", "Désignation", lab_w)] + [(f, h, w * k) for f, h, w in cols]
    tab = doc.add_table(rows=1, cols=len(heads))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    if excel:
        _xl_borders(tab)
    else:
        tab.style = "Table Grid"
    for i, (_, h, w) in enumerate(heads):
        c = tab.rows[0].cells[i]
        c.width = Cm(w)
        if excel:
            _xl_cell(c, h, XL_HEAD, size=9)
        else:
            c.text = ""
            _run(c.paragraphs[0], h, bold=True, size=8)
            _shade(c, "D9E2F3")
    ftxt = XL_FONT if excel else MATH
    for st in steps:
        cells = tab.add_row().cells
        for i, (k, _, w) in enumerate(heads):
            c = cells[i]
            c.width = Cm(w)
            if k == "label":
                if excel:
                    _xl_cell(c, st.label, XL_LABEL, size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
                else:
                    c.text = ""
                    _run(c.paragraphs[0], st.label, bold=True, size=8)
                if st.ref:
                    _run(c.add_paragraph(), st.ref, italic=True, size=6.5, color=GREY, font=XL_FONT if excel else None)
            elif k in ("sym", "num"):
                if excel:
                    _xl_cell(c, getattr(st, k), size=8, bold=False)
                else:
                    c.text = ""
                    _run(c.paragraphs[0], getattr(st, k), size=8, font=MATH)
            elif k == "res":
                if excel:
                    _xl_cell(c, st.res, XL_RES if st.res else None, size=8.5)
                else:
                    c.text = ""
                    _run(c.paragraphs[0], st.res, size=8, bold=True, color=BLUE)
                if st.ok is not None and st.cond:
                    _run(c.add_paragraph(), st.cond, size=6.5, italic=True, font=ftxt)
            else:  # verification
                txt = "" if st.ok is None else ("OK" if st.ok else "NON OK")
                if excel:
                    _xl_cell(c, txt, "FFC7CE" if st.ok is False else None, XL_RED, size=10)
                else:
                    c.text = ""
                    _run(c.paragraphs[0], txt, size=8, bold=True, color=GREEN if st.ok else RED)
                if st.ok is not None and st.cond and not COLS["res"]:
                    _run(c.add_paragraph(), st.cond, size=6.5, italic=True, font=ftxt)
    doc.add_paragraph()


_NUM = re.compile(r"^(\d+(?:\.\d+)*[a-z]?)\.?\s+")


def build_docx(tr: Trace, meta: dict, equations=True, selected=None, style=None, cover=True, toc=True,
               style_map=None, show_num=True, cols_map=None) -> bytes:
    """selected : set of section keys to print (None = all). style : 'latex' | 'tables' | 'excel' | 'plain'."""
    style = style or ("latex" if equations else "plain")
    EQUATIONS["on"] = style == "latex"
    EQUATIONS["num"] = show_num
    style_map = style_map or {}
    level_style = {0: style}
    cols_map = cols_map or {}
    level_cols = {0: dict(sym=True, num=show_num, res=True, ok=True)}
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    hdr = sec.header.paragraphs[0]
    _run(hdr, f"{meta.get('bureau', '')}   |   {meta.get('ouvrage', '')}   |   "
              f"{meta.get('titre_court', 'Note de calcul – Semelles isolées')}", size=8, color=GREY)
    ftr = sec.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(ftr, "Page ", size=8, color=GREY)
    _page_number(ftr)

    if cover:
        for _ in range(5):
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, "NOTE DE CALCUL", bold=True, size=26, color=BLUE)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, meta.get("sous_titre", "Fondations superficielles – Semelles isolées sous poteaux métalliques"), bold=True, size=14)
        doc.add_paragraph()
        tab = doc.add_table(rows=0, cols=2)
        tab.style = "Table Grid"
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in (("Projet / ouvrage", meta.get("ouvrage", "")), ("Maître d'ouvrage", meta.get("moa", "")),
                     ("Localisation", meta.get("lieu", "")), ("Bureau d'études", meta.get("bureau", "")),
                     ("Établi par", meta.get("ingenieur", "")), ("Date", meta.get("date", date.today().strftime("%d/%m/%Y"))),
                     ("Indice", meta.get("indice", "A")),
                     ("Règlements", meta.get("reglements", "RPA 2024 – BAEL 91 mod.99 – CBA 93 – DTR BC 2.331 – CCM 97 – Eurocode 7"))):
            c = tab.add_row().cells
            c[0].text, c[1].text = "", ""
            _run(c[0].paragraphs[0], k, bold=True)
            _shade(c[0], "D9E2F3")
            _run(c[1].paragraphs[0], str(v))
        doc.add_page_break()

    if toc:
        p = doc.add_paragraph()
        _run(p, "SOMMAIRE", bold=True, size=14, color=BLUE)
        p = doc.add_paragraph()
        r = p.add_run()
        for kind, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                r._r.append(el)
            elif text:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
                r._r.append(el)
            else:
                t = OxmlElement("w:t")
                t.text = "Clic droit → « Mettre à jour les champs » pour afficher le sommaire."
                r._r.append(t)
        doc.add_page_break()

    counters = [0, 0, 0, 0]
    skip_level = None
    first_l1 = True
    for s in tr.sections:
        # ---- selection: a skipped section hides its sub-sections
        if skip_level is not None:
            if s.level > skip_level:
                continue
            skip_level = None
        if selected is not None and s.key and s.key not in selected:
            skip_level = s.level
            continue
        if s.title:
            title = s.title
            m = _NUM.match(title)
            if m:  # automatic renumbering of the printed chapters
                lvl = min(s.level, 3)
                counters[lvl] += 1
                for k in range(lvl + 1, 4):
                    counters[k] = 0
                title = ".".join(str(counters[k]) for k in range(1, lvl + 1)) + " " + title[m.end():]
            if s.level == 1 and m and not first_l1 and s.key in ("semelles", "poteaux", "longrines", "recap"):
                doc.add_page_break()
            if s.level == 1:
                first_l1 = False
            h = doc.add_heading(title, level=min(s.level, 3))
            for r in h.runs:
                r.font.color.rgb = BLUE
        # ---- style of this chapter: its own choice, else inherited from its parent chapter
        cur = style_map.get(s.key) or level_style.get(s.level - 1, style)
        level_style[s.level] = cur
        for k in list(level_style):
            if k > s.level:
                del level_style[k]
        EQUATIONS["on"] = cur == "latex"
        ccur = dict(cols_map.get(s.key) or level_cols.get(s.level - 1, level_cols[0]))
        level_cols[s.level] = ccur
        for k in list(level_cols):
            if k > s.level:
                del level_cols[k]
        COLS.update(ccur)
        pending = []
        for it in s.items + [None]:
            if cur in ("tables", "excel") and isinstance(it, Step):
                pending.append(it)
                continue
            if pending:
                _grid(doc, pending, excel=cur == "excel")
                pending = []
            if it is None:
                break
            if isinstance(it, Step):
                _step(doc, it)
            elif isinstance(it, Para):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                if it.kind == "warn":
                    _run(p, "⚠ " + it.text, bold=True, color=RED)
                elif it.kind == "note":
                    _run(p, "Note : " + it.text, italic=True, color=GREY, size=9)
                elif it.kind == "bold":
                    _run(p, it.text, bold=True)
                elif it.kind == "sub":
                    r_ = _run(p, it.text, italic=True)
                    r_.underline = True
                else:
                    _run(p, it.text)
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif isinstance(it, Table):
                (_xl_table if cur == "excel" else _table)(doc, it)
            elif isinstance(it, Figure):
                doc.add_picture(BytesIO(it.png), width=Cm(it.width_cm))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if it.caption:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _run(p, it.caption, italic=True, size=9)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
