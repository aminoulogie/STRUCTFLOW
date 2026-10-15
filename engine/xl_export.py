"""Excel calculation sheet of each footing type, in the style of the user's own sheets
(Times New Roman bold, double borders, grey headers, yellow chosen values, green results, red OK).

Every table spans the same width (columns B … K): a table with fewer columns gets merged cells, so all the tables
line up for screenshots. Every result is a live Excel formula; with formulas=True a grey formula row is printed under
each table header (plus the explanatory note lines), with formulas=False only the short headers are kept.

Layout, top to bottom: characteristics → dimensioning (minimum vs chosen) → most critical loads of each combination →
weights → loads at the footing base and eccentricities (after adding the weights) → soil stresses X / Y →
overturning → sliding → uplift → punching → reinforcement // A and // B → conclusion."""
from __future__ import annotations

import math
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .footing import check_overturning, check_punching, check_sliding, check_uplift, d_values, rc_direction

TNR = "Times New Roman"
_d = Side(style="double", color="000000")
DBL = Border(left=_d, right=_d, top=_d, bottom=_d)
GREY = PatternFill("solid", fgColor="BFBFBF")
FGREY = PatternFill("solid", fgColor="E7E6E6")    # formula row
LIGHT = PatternFill("solid", fgColor="F2F2F2")
YEL = PatternFill("solid", fgColor="FFFF00")
GRN = PatternFill("solid", fgColor="00B050")
PINK = PatternFill("solid", fgColor="FFC7CE", bgColor="FFC7CE")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
NPHYS = 10                 # physical columns B … K
C0 = 2
WIDTHS = (2, 26, 15, 15, 26, 15, 15, 15, 15, 15, 15)  # A … K
OK = lambda cond: f'=IF({cond},"OK","NON OK")'  # noqa: E731
_UNIT = re.compile(r"\s*\(([^()=+−\-/·]{1,10})\)\s*$")


def split_head(t):
    """'Ntot = N + γG·W (kN)' → ('Ntot (kN)', 'Ntot = N + γG·W')."""
    if " = " not in t:
        return t, ""
    base, rest = t.split(" = ", 1)
    u = _UNIT.search(rest)
    unit = f" ({u.group(1)})" if u else ""
    formula = base + " = " + (rest[:u.start()] if u else rest)
    return base + unit, formula


def spans(n, weights=None):
    """Distribute the 10 physical columns over n logical columns (extra width to the first ones)."""
    if weights:
        assert sum(weights) == NPHYS and len(weights) == n
        return list(weights)
    base, extra = divmod(NPHYS, n)
    return [base + (1 if j < extra else 0) for j in range(n)]


class _Sheet:
    def __init__(self, ws, formulas=True):
        self.ws = ws
        self.formulas = formulas
        self.r = 1
        self.cur = spans(NPHYS)
        self.lay = {}                 # row → spans used on that row
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ---- layout
    def cols(self, n, weights=None):
        self.cur = spans(n, weights)

    def assign(self, *rows):
        for r in rows:
            self.lay[r] = self.cur

    def _phys(self, r, j):
        sp = self.lay.setdefault(r, self.cur)
        c0 = C0 + sum(sp[:j])
        return c0, c0 + sp[j] - 1

    # ---- cells (j = logical column)
    def cell(self, r, j, v, fill=None, fmt=None, color="000000", bold=True, size=11, align=CEN):
        c0, c1 = self._phys(r, j)
        if c1 > c0:
            self.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        for c in range(c0, c1 + 1):
            x = self.ws.cell(r, c)
            x.border = DBL
            if fill:
                x.fill = fill
        x = self.ws.cell(r, c0)
        x.value = v
        x.font = Font(name=TNR, size=size, bold=bold, color=color)
        x.alignment = align
        if fmt:
            x.number_format = fmt
        return x

    def ref(self, r, j):
        c0, _ = self._phys(r, j)
        return f"${get_column_letter(c0)}${r}"

    def okcell(self, r, j, cond):
        return self.cell(r, j, OK(cond), color="FF0000", size=12)

    # ---- full-width rows
    def band(self, r, v, fill=GREY, size=12, bold=True, align=CEN, border=True, color="000000"):
        self.ws.merge_cells(start_row=r, start_column=C0, end_row=r, end_column=C0 + NPHYS - 1)
        for c in range(C0, C0 + NPHYS):
            x = self.ws.cell(r, c)
            if border:
                x.border = DBL
            if fill:
                x.fill = fill
        x = self.ws.cell(r, C0)
        x.value = v
        x.font = Font(name=TNR, size=size, bold=bold, color=color)
        x.alignment = align
        return x

    def section(self, title):
        self.r += 1
        self.band(self.r, title, GREY, 12)
        self.ws.row_dimensions[self.r].height = 20
        self.r += 1

    def note(self, text, force=False):
        if not (self.formulas or force):
            return
        self.band(self.r, text, None, 10, bold=False, align=LEFT, border=False)
        self.ws.row_dimensions[self.r].height = 15 * max(1, math.ceil(len(text) / 140))
        self.r += 1

    def header(self, heads, weights=None, height=None):
        """New table: sets its column layout, prints the short headers and (formulas on) the formula row."""
        self.cols(len(heads), weights)
        parts = [split_head(h) for h in heads]
        for j, (short, _) in enumerate(parts):
            self.cell(self.r, j, short, GREY, size=10)
        self.ws.row_dimensions[self.r].height = height or (30 if max(len(p[0]) for p in parts) > 14 else 20)
        self.r += 1
        if self.formulas and any(f for _, f in parts):
            for j, (_, f) in enumerate(parts):
                self.cell(self.r, j, f, FGREY, size=9, bold=False)
            self.ws.row_dimensions[self.r].height = 30 if max(len(f) for _, f in parts) > 18 else 18
            self.r += 1


def _pick_rows(res):
    """(combo, joint) pairs to print: per combination the most critical joint, plus every pair that governs a check
    in the application, so that the sheet reproduces the application's maxima."""
    combos = [c for c in res.combos if c.ctype in ("ELU", "ELS", "ACC")]
    names = list(dict.fromkeys(c.name for c in combos))
    score = {}
    ev = res.ev
    for lst in (ev["bear"], ev["ovt"], ev["sld"], res.punch):
        for r in lst:
            v = r.get("ratio")
            if v is not None and not math.isinf(v):
                score[(r["combo"], r["joint"])] = max(score.get((r["combo"], r["joint"]), 0.0), v)
    keep = set()
    for n in names:
        js = [c for c in combos if c.name == n]
        best = max(js, key=lambda c: (score.get((c.name, c.joint), 0.0), abs(c.N)))
        keep.add((best.name, best.joint))
    gov = lambda lst: max((r for r in lst if r.get("ratio") is not None), key=lambda r: r["ratio"], default=None)  # noqa: E731
    for lst in (ev["bear"], ev["ovt"], ev["sld"], ev["upl"], res.punch):
        g = gov(lst)
        if g:
            keep.add((g["combo"], g["joint"]))
    for key in ("X", "Y"):
        for r in res.rc[key]["cand"].values():
            if r:
                keep.add((r["combo"], r["joint"]))
    out = [c for c in combos if (c.name, c.joint) in keep]
    multi = len({c.joint for c in out}) > 1 or len(out) != len(names)
    return out, multi


def footing_sheet(ws, prj, res, formulas=True):
    S = _Sheet(ws, formulas)
    m, soil, o, cd = prj.mat, prj.soil, prj.opts, prj.codes
    g, gi = res.geo, res.gi
    caps = res.ev["caps"]
    rows, multi = _pick_rows(res)
    lab = lambda c: f"{c.name} (appui {c.joint})" if multi else c.name  # noqa: E731

    # ------------------------------------------------------------------ title
    S.band(1, f"FONDATIONS : SEMELLES ISOLÉES – SEMELLE {gi.name}", GREY, 14)
    S.ws.row_dimensions[1].height = 24
    S.band(2, f"{prj.meta.get('ouvrage', '')} – appuis : {', '.join(str(j) for j in gi.joints)} – pied de poteau : "
              f"{gi.base.lower()}", None, 10, bold=False, border=False)
    S.r = 3

    # ------------------------------------------------------------------ 1. characteristics
    S.section("1. CARACTÉRISTIQUES DES MATÉRIAUX, DU SOL ET DU POTEAU")
    sadm = soil.sigma_adm / 100 if soil.method == "adm" else caps["q_els"] / 100
    S.header(["Fe (MPa)", "FC28 (MPa)", "ft28 = 0,6 + 0,06·fc28 (MPa)", "γ béton (kN/m³)", "Enrobage c (m)",
              "Ancrage D (m)", "σ sol (bar)", "γ sol (kN/m³)", "φ' (°)", "c' (kPa)"])
    r = S.r
    for j, v in enumerate([m.fe, m.fc28, None, m.gamma_c, m.cover, g.D, sadm, soil.gamma_r, soil.phi_s, soil.c_s]):
        S.cell(r, j, v, fmt="0.00")
    S.cell(r, 2, f"=0.6+0.06*{S.ref(r, 1)}", fmt="0.00")
    FE, FC, FT, GC, COV, D, _, GR, PHI, CS = (S.ref(r, j) for j in range(10))
    S.r += 1
    sb = m.sigma_s_bar
    S.header(["q ELS (bar)", "q ELU (bar)", "q ACC (bar)", f"σs ELU = fe/{m.gamma_s:g} (MPa)",
              f"σs ACC = fe/{m.gamma_s_acc:g} (MPa)", "σ̄s ELS (MPa)", f"fbu ELU = 0,85·fc28/{m.theta * m.gamma_b:g} (MPa)",
              f"fbu ACC = 0,85·fc28/{m.theta * m.gamma_b_acc:g} (MPa)", "σ̄bc = 0,6·fc28 (MPa)", "φ supposé (mm)"])
    r = S.r
    S.cell(r, 0, caps["q_els"] / 100, fmt="0.00")
    S.cell(r, 1, caps["q_elu"] / 100, fmt="0.00")
    S.cell(r, 2, caps["q_acc"] / 100, fmt="0.00")
    S.cell(r, 3, f"={FE}/{m.gamma_s}", fmt="0.00")
    S.cell(r, 4, f"={FE}/{m.gamma_s_acc}", fmt="0.00")
    S.cell(r, 5, sb if sb else "non limitée", fmt="0.00")
    S.cell(r, 6, f"=0.85*{FC}/({m.theta}*{m.gamma_b})", fmt="0.00")
    S.cell(r, 7, f"=0.85*{FC}/({m.theta}*{m.gamma_b_acc})", fmt="0.00")
    S.cell(r, 8, f"=0.6*{FC}", fmt="0.00")
    S.cell(r, 9, o.phi_assumed * 1000, fmt="0")
    Q = {"ELS": S.ref(r, 0), "ELU": S.ref(r, 1), "ACC": S.ref(r, 2)}
    SS = {False: S.ref(r, 3), True: S.ref(r, 4)}
    SSB = S.ref(r, 5)
    FBU = {False: S.ref(r, 6), True: S.ref(r, 7)}
    PHID = S.ref(r, 9)
    S.r += 1
    S.header(["Poteau / fût a (m)", "Poteau / fût b (m)", "h0 arase / TN (m)", "Hp = D − H + h0 (m)",
              "H = D + h0 (m)" if g.z_app is None else "H = D + z appuis (m)", "μl ELU", "μl ACC", "γb ELU", "γb ACC",
              "Fissuration"])
    rch = S.r
    S.cell(rch, 0, g.ap, fmt="0.00")
    S.cell(rch, 1, g.bp, fmt="0.00")
    S.cell(rch, 2, g.h_above, fmt="0.00")
    AP, BP, H0 = S.ref(rch, 0), S.ref(rch, 1), S.ref(rch, 2)
    HP, HL = S.ref(rch, 3), S.ref(rch, 4)
    S.cell(rch, 5, round(m.mu_l(False), 4), fmt="0.000")
    S.cell(rch, 6, round(m.mu_l(True), 4), fmt="0.000")
    S.cell(rch, 7, m.gamma_b, fmt="0.00")
    S.cell(rch, 8, m.gamma_b_acc, fmt="0.00")
    S.cell(rch, 9, m.cracking)
    MUL = {False: S.ref(rch, 5), True: S.ref(rch, 6)}
    GB = {False: S.ref(rch, 7), True: S.ref(rch, 8)}
    S.r += 1

    # ------------------------------------------------------------------ 2. dimensioning
    S.section("2. DIMENSIONNEMENT DE LA SEMELLE")
    S.header(["", "A (m)", "B (m)", "H (m)", "dx (m)", "dy (m)"], [3, 1, 1, 1, 2, 2])
    rmin, rret, rchk = S.r, S.r + 1, S.r + 2
    S.assign(rmin, rret, rchk)
    S.cell(rret, 0, "Retenu (choisi)", LIGHT)
    S.cell(rret, 1, g.A, YEL, "0.00")
    S.cell(rret, 2, g.B, YEL, "0.00")
    S.cell(rret, 3, g.h, YEL, "0.00")
    A, B, H = S.ref(rret, 1), S.ref(rret, 2), S.ref(rret, 3)
    dl, ds = f"={H}-{COV}-{PHID}/2000", f"={H}-{COV}-1.5*{PHID}/1000"
    S.cell(rret, 4, dl if g.A >= g.B else ds, fmt="0.000")
    S.cell(rret, 5, ds if g.A >= g.B else dl, fmt="0.000")
    DX, DY = S.ref(rret, 4), S.ref(rret, 5)
    S.cell(rch, 3, f"={D}-{H}+{H0}", fmt="0.00")
    S.cell(rch, 4, f"={D}+{H0}" if g.z_app is None else f"={D}+({g.z_app})", fmt="0.00")
    S.r = rchk + 1
    S.note("Surface minimale (contrainte moyenne de service) : Nser / (A·B) ≤ q_ELS  ⇒  S = A·B ≥ Nser,max / q_ELS ; "
           "homothétie avec le poteau : A = √(S·a/b), B = √(S·b/a). Rigidité : d ≥ max((A − a)/4 ; (B − b)/4) ; "
           "H ≥ d + c + 1,5·φ (arrondi à 5 cm).")
    S.header(["Nser,max (kN)", "S min = Nser/q_ELS (m²)", "A min = √(S·a/b) (m)", "B min = √(S·b/a) (m)",
              "d min = max((A−a)/4 ; (B−b)/4) (m)", "H min = d + c + 1,5φ (m)"], [2, 2, 1, 1, 2, 2])
    rdv = S.r
    S.assign(rdv)
    S.r += 1

    # ------------------------------------------------------------------ 3. loads
    S.section("3. SOLLICITATIONS LES PLUS DÉFAVORABLES DE CHAQUE COMBINAISON (réactions d'appui, repère global)")
    S.header(["Combinaison", "Type", "Appui", "N (kN)", "Mx (kN.m)", "My (kN.m)", "Vx (kN)", "Vy (kN)"], [3, 1, 1, 1, 1, 1, 1, 1])
    L = {}
    for c in rows:
        r = S.r
        S.cell(r, 0, c.name, LIGHT)
        S.cell(r, 1, c.ctype)
        S.cell(r, 2, c.joint)
        for j, v in enumerate((c.N, c.Mx, c.My, c.Vx, c.Vy)):
            S.cell(r, 3 + j, round(v, 3), GRN if j == 0 else None, "0.00")
        L[(c.name, c.joint)] = dict(N=S.ref(r, 3), Mx=S.ref(r, 4), My=S.ref(r, 5), Vx=S.ref(r, 6), Vy=S.ref(r, 7))
        S.r += 1
    els_N = [L[(c.name, c.joint)]["N"] for c in rows if c.ctype == "ELS"]
    nser = f"MAX({','.join(els_N)})" if els_N else f"MAX({','.join(L[(c.name, c.joint)]['N'] for c in rows)})/1.35"
    S.cell(rdv, 0, f"={nser}", fmt="0.00")
    S.cell(rdv, 1, f"={S.ref(rdv, 0)}/({Q['ELS']}*100)", fmt="0.000")
    S.cell(rdv, 2, f"=SQRT({S.ref(rdv, 1)}*{AP}/{BP})", GRN, "0.00")
    S.cell(rdv, 3, f"=SQRT({S.ref(rdv, 1)}*{BP}/{AP})", GRN, "0.00")
    S.cell(rdv, 4, f"=MAX(({A}-{AP})/4,({B}-{BP})/4)", fmt="0.000")
    S.cell(rdv, 5, f"=MAX({o.h_min},CEILING({S.ref(rdv, 4)}+{COV}+1.5*{PHID}/1000,0.05))", GRN, "0.00")
    S.cell(rmin, 0, "Minimum requis", LIGHT)
    for j, src in ((1, 2), (2, 3), (3, 5), (4, 4), (5, 4)):
        S.cell(rmin, j, f"={S.ref(rdv, src)}", fmt="0.000" if j >= 4 else "0.00")
    S.cell(rchk, 0, "Vérification", LIGHT)
    for j in range(1, 6):
        S.okcell(rchk, j, f"{S.ref(rret, j)}>={S.ref(rmin, j)}-0.0001")

    # ------------------------------------------------------------------ 4. weights
    S.section("4. POIDS DES ÉLÉMENTS")
    S.header(["Élément", "Formule", "Poids (kN)"] if formulas else ["Élément", "Poids (kN)"],
             [3, 5, 2] if formulas else [6, 4])
    W = {}
    for key, labw, form, f in (
            ("Wf", "Poids de la semelle", "Wf = γb·A·B·H", f"={GC}*{A}*{B}*{H}"),
            ("Wp", "Poids du fût (amorce)", "Wp = γb·a·b·Hp", f"={GC}*{AP}*{BP}*{HP}"),
            ("Ws", "Poids des terres", "Ws = γsol·(A·B − a·b)·(D − H)", f"={GR}*({A}*{B}-{AP}*{BP})*MAX({D}-{H},0)"),
            ("W", "POIDS TOTAL", "W = Wf + Wp + Ws", None)):
        r = S.r
        S.cell(r, 0, labw, GREY if key == "W" else LIGHT)
        if key == "W":
            f = f"={W['Wf']}+{W['Wp']}+{W['Ws']}"
        jv = 2 if formulas else 1
        if formulas:
            S.cell(r, 1, form, size=11, bold=False)
        S.cell(r, jv, f, GRN if key == "W" else None, "0.00")
        W[key] = S.ref(r, jv)
        S.r += 1

    # ------------------------------------------------------------------ 5. loads at the base
    S.section("5. EFFORTS À LA SOUS-FACE DE LA SEMELLE ET EXCENTRICITÉS (poids des éléments inclus)")
    S.note("Le poids des éléments est ajouté AVANT le calcul de l'excentricité. γG = 1,35 à l'ELU, 1,00 à l'ELS et "
           "en situation sismique. Les moments sont transportés au niveau d'assise par le bras de levier H.")
    S.header(["Combinaison", "γG", "Ntot = N + γG·W (kN)", "My,b = My + Vx·H (kN.m)", "Mx,b = Mx − Vy·H (kN.m)",
              "ex = |My,b| / Ntot (m)", "ey = |Mx,b| / Ntot (m)"], [3, 1, 2, 1, 1, 1, 1])
    BASE = {}
    for c in rows:
        r, l_ = S.r, L[(c.name, c.joint)]
        S.cell(r, 0, lab(c), LIGHT)
        S.cell(r, 1, c.g_unf, fmt="0.00")
        S.cell(r, 2, f"={l_['N']}+{S.ref(r, 1)}*{W['W']}", GRN, "0.00")
        S.cell(r, 3, f"={l_['My']}+{l_['Vx']}*{HL}", fmt="0.00")
        S.cell(r, 4, f"={l_['Mx']}-{l_['Vy']}*{HL}", fmt="0.00")
        nt = S.ref(r, 2)
        S.cell(r, 5, f'=IF({nt}>0,ABS({S.ref(r, 3)})/{nt},"N ≤ 0")', fmt="0.000")
        S.cell(r, 6, f'=IF({nt}>0,ABS({S.ref(r, 4)})/{nt},"N ≤ 0")', fmt="0.000")
        BASE[(c.name, c.joint)] = dict(Nt=nt, Myb=S.ref(r, 3), Mxb=S.ref(r, 4), ex=S.ref(r, 5), ey=S.ref(r, 6))
        S.r += 1

    # ------------------------------------------------------------------ 6. soil stresses
    for sens, Lr, Wr, ek, ls in (("X", A, B, "ex", "A"), ("Y", B, A, "ey", "B")):
        wl = "B" if ls == "A" else "A"
        S.section(f'6.{1 if sens == "X" else 2} VÉRIFICATION DES CONTRAINTES DU SOL "Sens_{sens}" (DTR BC 2.331 – règle des 3/4)')
        S.note(f"e ≤ {ls}/6 : diagramme trapézoïdal ; e > {ls}/6 : diagramme triangulaire (σmin = 0) ; "
               "contrainte de référence σ3/4 = (3σmax + σmin)/4 ≤ q.")
        S.header(["Combinaison", "Ntot (kN)", f"{ek} (m)", "Diagramme",
                  f"σmax = Ntot/(A·B)·(1 + 6e/{ls}) ou 2Ntot/(3{wl}({ls}/2 − e)) (bar)",
                  f"σmin = Ntot/(A·B)·(1 − 6e/{ls}) ou 0 (bar)", "σ 3/4 = (3σmax + σmin)/4 (bar)", "q (bar)",
                  "Taux = σ3/4 / q", "Vérification"])
        for c in rows:
            r, b_ = S.r, BASE[(c.name, c.joint)]
            N_, e_ = b_["Nt"], b_[ek]
            S.cell(r, 0, lab(c), LIGHT)
            S.cell(r, 1, f"={N_}", GRN, "0.00")
            S.cell(r, 2, f"={e_}", fmt="0.000")
            S.cell(r, 3, f'=IF({N_}<=0,"N ≤ 0 : soulèvement",IF({e_}<={Lr}/6+0.000001,"e < {ls}/6 => Diagramme Trapézoïdal",'
                         f'IF({e_}<{Lr}/2,"e > {ls}/6 => Diagramme triangulaire","e ≥ {ls}/2 : instable")))', size=10)
            S.cell(r, 4, f"=IF({N_}<=0,0,IF({e_}<={Lr}/6+0.000001,{N_}/({Lr}*{Wr})*(1+6*{e_}/{Lr}),"
                         f"IF({e_}<{Lr}/2,2*{N_}/(3*{Wr}*({Lr}/2-{e_})),1E9)))/100", fmt="0.00")
            S.cell(r, 5, f"=IF({N_}<=0,0,IF({e_}<={Lr}/6+0.000001,{N_}/({Lr}*{Wr})*(1-6*{e_}/{Lr}),0))/100", fmt="0.00")
            S.cell(r, 6, f"=(3*{S.ref(r, 4)}+{S.ref(r, 5)})/4", GRN, "0.00")
            S.cell(r, 7, f"={Q[c.ctype]}", fmt="0.00")
            S.cell(r, 8, f"={S.ref(r, 6)}/{S.ref(r, 7)}", fmt="0%")
            S.okcell(r, 9, f"{S.ref(r, 6)}<={S.ref(r, 7)}+0.00001")
            S.r += 1

    # ------------------------------------------------------------------ 7. overturning
    ovt_rows = [(c, c.ctype == "ACC") for c in rows if check_overturning(g, c, o, cd)]
    if ovt_rows:
        S.section("7. VÉRIFICATION AU RENVERSEMENT")
        S.note("Moment stabilisant calculé au niveau bas de la fondation par rapport au bord extrême. Séisme : Fs ≥ 1,3 "
               f"avec le poids total (RPA 2024 §4.5.1 et §5.5) ; ELS / ELU : Fs ≥ {o.fs_overturn:g} (règle de l'art).")
        S.header(["Combinaison", "γG", "Nt = N + γG·W (kN)", "Fs x = Nt·(A/2) / |My,b|", "Fs y = Nt·(B/2) / |Mx,b|",
                  "Fs requis", "Vérification"], [3, 1, 2, 1, 1, 1, 1])
        for c, rpa in ovt_rows:
            r, b_, l_ = S.r, BASE[(c.name, c.joint)], L[(c.name, c.joint)]
            S.cell(r, 0, lab(c), LIGHT)
            S.cell(r, 1, 1.0 if rpa else c.g_fav, fmt="0.00")
            S.cell(r, 2, f"={l_['N']}+{S.ref(r, 1)}*{W['W']}", GRN, "0.00")
            nt = S.ref(r, 2)
            S.cell(r, 3, f"=IF({nt}<=0,0,IF(ABS({b_['Myb']})<0.000001,999,{nt}*{A}/2/ABS({b_['Myb']})))", fmt="0.00")
            S.cell(r, 4, f"=IF({nt}<=0,0,IF(ABS({b_['Mxb']})<0.000001,999,{nt}*{B}/2/ABS({b_['Mxb']})))", fmt="0.00")
            S.cell(r, 5, 1.3 if rpa else o.fs_overturn, fmt="0.00")
            S.okcell(r, 6, f"MIN({S.ref(r, 3)},{S.ref(r, 4)})>={S.ref(r, 5)}-0.00001")
            S.r += 1

    # ------------------------------------------------------------------ 8. sliding
    sld_rows = [c for c in rows if check_sliding(g, c, soil, o)]
    if sld_rows:
        S.section("8. VÉRIFICATION AU GLISSEMENT")
        S.note("A' = Lx'·Ly' = surface comprimée (L' = L si e ≤ L/6, sinon 3·(L/2 − e)). Fs ≥ 1,25 en situation sismique "
               f"(RPA 2024 §5.5 et §10.1.5) ; ≥ {o.fs_slide:g} en ELS / ELU (DTR BC 2.33.1).")
        S.header(["Combinaison", "H = √(Vx² + Vy²) (kN)", "Nt = N + γG·W (kN)", "Lx' (m)", "Ly' (m)",
                  "R = Nt·tanφ' + c'·A' (kN)", "Fs = R / H", "Fs requis", "Vérification"], [2, 1, 1, 1, 1, 1, 1, 1, 1])
        for c in sld_rows:
            r, b_, l_ = S.r, BASE[(c.name, c.joint)], L[(c.name, c.joint)]
            S.cell(r, 0, lab(c), LIGHT)
            S.cell(r, 1, f"=SQRT({l_['Vx']}^2+{l_['Vy']}^2)", fmt="0.00")
            S.cell(r, 2, f"={l_['N']}+{c.g_fav}*{W['W']}", GRN, "0.00")
            nt = S.ref(r, 2)
            ex_, ey_ = f"ABS({b_['Myb']})/{nt}", f"ABS({b_['Mxb']})/{nt}"
            S.cell(r, 3, f"=IF({nt}<=0,0,IF({ex_}<={A}/6+0.000001,{A},MIN(3*({A}/2-{ex_}),{A})))", fmt="0.000")
            S.cell(r, 4, f"=IF({nt}<=0,0,IF({ey_}<={B}/6+0.000001,{B},MIN(3*({B}/2-{ey_}),{B})))", fmt="0.000")
            S.cell(r, 5, f"=MAX({nt},0)*TAN(RADIANS({PHI}))+{CS}*{S.ref(r, 3)}*{S.ref(r, 4)}", fmt="0.00")
            S.cell(r, 6, f"={S.ref(r, 5)}/{S.ref(r, 1)}", fmt="0.00")
            S.cell(r, 7, o.fs_slide_acc if c.ctype == "ACC" else o.fs_slide, fmt="0.00")
            S.okcell(r, 8, f"{S.ref(r, 6)}>={S.ref(r, 7)}-0.00001")
            S.r += 1

    # ------------------------------------------------------------------ 9. uplift
    has_upl = any(c.ctype == "UPL" for c in res.combos)
    upl_rows = [c for c in rows if check_uplift(g, c, o, cd, has_upl)]
    S.section("9. VÉRIFICATION AU SOULÈVEMENT")
    if upl_rows:
        S.note(f"Fs requis : {o.fs_uplift_els:g} (ELS) ou {o.fs_uplift_elu:g} (ELU, séisme) – règle de l'art.")
        S.header(["Combinaison", "T = −N (kN)", "W stab = γG·W (kN)", "Fs = γG·W / T", "Fs requis", "Vérification"],
                 [3, 2, 2, 1, 1, 1])
        for c in upl_rows:
            r, l_ = S.r, L[(c.name, c.joint)]
            S.cell(r, 0, lab(c), LIGHT)
            S.cell(r, 1, f"=-{l_['N']}", fmt="0.00")
            S.cell(r, 2, f"={c.g_fav}*{W['W']}", GRN, "0.00")
            S.cell(r, 3, f"={S.ref(r, 2)}/{S.ref(r, 1)}", fmt="0.00")
            S.cell(r, 4, o.fs_uplift_els if c.ctype == "ELS" else o.fs_uplift_elu, fmt="0.00")
            S.okcell(r, 5, f"{S.ref(r, 3)}>={S.ref(r, 4)}-0.00001")
            S.r += 1
    else:
        S.note("Aucune combinaison ne donne de traction au pied du poteau (N > 0) : vérification sans objet.", force=True)

    # ------------------------------------------------------------------ 10. punching
    pu_rows = [c for c in rows if check_punching(g, c, m)]
    if pu_rows:
        S.section("10. VÉRIFICATION AU POINÇONNEMENT (BAEL A.5.2.4)")
        S.header(["Combinaison", "Nu = N + γG·Wp (kN)", "P'u = Nu·[1 − (a + 2H)(b + 2H)/(A·B)] (kN)",
                  "uc = 2·(a + b + 2H) (m)", "γb", "P'lim = 0,045·uc·H·fc28/γb (kN)", "Taux = P'u / P'lim", "Vérification"],
                 [3, 1, 1, 1, 1, 1, 1, 1])
        for c in pu_rows:
            r, l_ = S.r, L[(c.name, c.joint)]
            S.cell(r, 0, lab(c), LIGHT)
            S.cell(r, 1, f"={l_['N']}+{c.g_unf}*{W['Wp']}", GRN, "0.00")
            nu = S.ref(r, 1)
            S.cell(r, 2, f"=IF(AND({AP}+2*{H}>={A},{BP}+2*{H}>={B}),0,{nu}*(1-MIN({AP}+2*{H},{A})*MIN({BP}+2*{H},{B})/({A}*{B})))",
                   fmt="0.00")
            S.cell(r, 3, f"=2*({AP}+{BP}+2*{H})", fmt="0.000")
            S.cell(r, 4, f"={GB[c.acc]}", fmt="0.00")
            S.cell(r, 5, f"=0.045*{S.ref(r, 3)}*{H}*{FC}/{S.ref(r, 4)}*1000", fmt="0.00")
            S.cell(r, 6, f"={S.ref(r, 2)}/{S.ref(r, 5)}", fmt="0%")
            S.okcell(r, 7, f"{S.ref(r, 2)}<={S.ref(r, 5)}+0.00001")
            S.r += 1

    # ------------------------------------------------------------------ 11. reinforcement
    concl = []
    dvals = d_values(g, m, o)
    for k, (key, Lr, Wr, lr, dref, Ls, Ws, ls) in enumerate((("X", A, B, AP, DX, "A", "B", "a"),
                                                              ("Y", B, A, BP, DY, "B", "A", "b"))):
        S.section(f"11.{k + 1} FERRAILLAGE DE LA SEMELLE – NAPPE INFÉRIEURE // {Ls} ({'x' if key == 'X' else 'y'})")
        S.note(f"Méthode des bielles si e₀ = |M|/Nu ≤ {Ls}/6. Méthode des moments à la section S1 (l₁ = {Ls}/2 − 0,35·{ls}) : "
               f"M₁ = {Ws}·∫₀^l₁ σ(s)·(l₁ − s)·ds − w·{Ws}·l₁²/2, w = γG·(γb·H + γsol·(D − H)). ELU / ACC : z = d·(1 − 0,4α), "
               "α = 1,25(1 − √(1 − 2μ)). ELS : section fissurée n = 15, α³ − 3α² − 90μs·α + 90μs = 0, z = d·(1 − α/3).")
        S.header(["Combinaison", "Nu = N + γG·Wp (kN)", "e₀ = |M| / Nu (m)", f"As bielles = N'·({Ls} − {ls})/(8·d·σs) (cm²)",
                  f"x comprimée = {Ls} ou 3({Ls}/2 − e) (m)", "σmax (kPa)", "M₁ (S1) (kN.m)", "μ = M₁/(b·d²·fbu ou σ̄s)",
                  "As moments = M₁/(z·σs) (cm²)", "As = max (cm²)"])
        r_first = S.r
        l1 = f"({Lr}/2-0.35*{lr})"
        mk = "Myb" if key == "X" else "Mxb"
        for c in rows:
            if c.ctype == "ELS" and m.sigma_s_bar is None:
                continue
            if not rc_direction(g, c, key, dvals[0 if key == "X" else 1], m):
                continue
            r, b_, l_ = S.r, BASE[(c.name, c.joint)], L[(c.name, c.joint)]
            els = c.ctype == "ELS"
            ss = SSB if els else SS[c.acc]
            S.cell(r, 0, lab(c), LIGHT)
            S.cell(r, 1, f"={l_['N']}+{c.g_unf}*{W['Wp']}", fmt="0.00")
            nu = S.ref(r, 1)
            S.cell(r, 2, f'=IF({nu}>0,ABS({b_[mk]})/{nu},"—")', fmt="0.000")
            e0 = S.ref(r, 2)
            S.cell(r, 3, f'=IF({nu}<=0,"—",IF({e0}<={Lr}/6+0.000001,10*{nu}*(1+3*{e0}/{Lr})*({Lr}-{lr})/(8*{dref}*{ss}),"—"))',
                   fmt="0.00")
            nt = b_["Nt"]
            et = f"ABS({b_[mk]})/{nt}"
            S.cell(r, 4, f"=IF({nt}<=0,0,IF({et}<={Lr}/6+0.000001,{Lr},3*({Lr}/2-{et})))", fmt="0.000")
            x = S.ref(r, 4)
            S.cell(r, 5, f"=IF({nt}<=0,0,IF({et}<={Lr}/6+0.000001,{nt}/({Lr}*{Wr})*(1+6*{et}/{Lr}),2*{nt}/({Wr}*{x})))",
                   fmt="0.0")
            sm = S.ref(r, 5)
            kk = f"IF({et}<={Lr}/6+0.000001,12*{nt}*{et}/({Wr}*{Lr}^3),{sm}/{x})"
            s_ = f"MIN({l1},{x})"
            w = f"{c.g_unf}*({GC}*{H}+{GR}*MAX({D}-{H},0))"
            S.cell(r, 6, f"=IF({nt}<=0,0,{Wr}*({sm}*({l1}*{s_}-{s_}^2/2)-{kk}*({l1}*{s_}^2/2-{s_}^3/3)))-{w}*{Wr}*{l1}^2/2",
                   fmt="0.00")
            m1 = S.ref(r, 6)
            if els:
                S.cell(r, 7, f"=MAX({m1},0)/({Wr}*{dref}^2*{ss}*1000)", fmt="0.0000")
                mu = S.ref(r, 7)
                p = f"(3+90*{mu})"
                alpha = f"(1+2*SQRT({p}/3)*COS(ACOS((3/{p})*SQRT(3/{p}))/3-2*PI()/3))"
                S.cell(r, 8, f"=IF({m1}<=0,0,10*{m1}/({dref}*(1-{alpha}/3)*{ss}))", fmt="0.00")
            else:
                S.cell(r, 7, f"=MAX({m1},0)/({Wr}*{dref}^2*{FBU[c.acc]}*1000)", fmt="0.0000")
                mu = S.ref(r, 7)
                S.ws.conditional_formatting.add(mu.replace("$", ""), FormulaRule(formula=[f"{mu}>{MUL[c.acc]}"], fill=PINK))
                S.cell(r, 8, f'=IF({m1}<=0,0,IF({mu}>=0.5,"μ ≥ 0,5",10*{m1}/({dref}*(1-0.4*1.25*(1-SQRT(1-2*{mu})))*{ss})))',
                       fmt="0.00")
            S.cell(r, 9, f"=MAX(N({S.ref(r, 3)}),N({S.ref(r, 8)}))", GRN, "0.00")
            S.r += 1
        r_last = S.r - 1
        rc = res.rc[key]
        bars = rc["bars"]
        # summary – same width: Désignation | (Formule) | Valeur | Vérification
        S.r += 1
        S.header(["Désignation", "Formule", "Valeur", "Vérification"] if formulas else ["Désignation", "Valeur", "Vérification"],
                 [3, 4, 2, 1] if formulas else [6, 2, 2])
        summ = {}

        def line(tag, labl, form, f, fill=None, fmt="0.00", chk=None):
            r = S.r
            S.cell(r, 0, labl, LIGHT, align=LEFT, size=10)
            jv = 2 if formulas else 1
            if formulas:
                S.cell(r, 1, form, size=10, bold=False)
            S.cell(r, jv, f, fill, fmt)
            summ[tag] = S.ref(r, jv)
            if chk:
                S.okcell(r, jv + 1, chk.format(**summ))
            else:
                S.cell(r, jv + 1, "")
            S.r += 1

        line("Asc", "As max (bielles / moments)", "max des combinaisons",
             f"=MAX({S.ref(r_first, 9)}:{S.ref(r_last, 9)})" if r_last >= r_first else 0)
        line("Amin", "As min – non-fragilité", f"0,23·{Ws}·d·ft28/fe", f"=0.23*{Wr}*{dref}*{FT}/{FE}*10000")
        line("Asr", "As requis", "max(As ; As min)", f"=MAX({summ['Asc']},{summ['Amin']})", GRN)
        line("phi", "Diamètre retenu φ (mm)", "choix", bars["phi"], YEL, "0")
        line("s", "Espacement retenu s (m)", "choix", round(bars["s"], 3), YEL, "0.00")
        line("n", "Nombre de barres", f"n = ⌈({Ws} − 2c)/s⌉ + 1",
             f"=MAX(CEILING(({Wr}-2*{COV})/{summ['s']}-0.000000001,1)+1,2)", None, "0")
        line("Asf", "As fourni", "n·π·φ²/4", f"={summ['n']}*PI()*{summ['phi']}^2/400", GRN, "0.00", "{Asf}>={Asr}-0.0001")
        line("smax", "Espacement maximal", {"FTP": "min(1,5·H ; 20 cm)", "FP": "min(2·H ; 25 cm)"}.get(m.cracking, "min(3·H ; 33 cm)"),
             {"FTP": f"=MIN(1.5*{H},0.2)", "FP": f"=MIN(2*{H},0.25)"}.get(m.cracking, f"=MIN(3*{H},0.33)"), None, "0.00",
             "{s}<={smax}+0.0001")
        S.cols(2, [3, 7])
        r = S.r
        S.cell(r, 0, f"FERRAILLAGE // {Ls}", GREY)
        S.cell(r, 1, f'={summ["n"]}&" HA"&{summ["phi"]}&"  e = "&TEXT({summ["s"]}*100,"0")&" cm  (As = "'
                     f'&TEXT({summ["Asf"]},"0.00")&" cm²)"', YEL, size=12)
        S.r += 1
        if rc.get("top"):
            S.note(f"Nappe supérieure nécessaire (M₁ < 0 : soulèvement de la console) : As,sup = {rc['As_top']:.2f} cm² "
                   + (f"→ {rc['bars_top']['n']} HA{rc['bars_top']['phi']}" if rc.get("bars_top") else ""), force=True)
        concl.append(f'{summ["n"]}&" HA"&{summ["phi"]}&" e = "&TEXT({summ["s"]}*100,"0")&" cm"')

    # ------------------------------------------------------------------ conclusion
    S.r += 1
    r = S.r
    rng = f"$B$4:$K${r - 1}"
    S.cols(2, [3, 7])
    S.cell(r, 0, "CONCLUSION", GREY, size=12)
    S.cell(r, 1, f'=IF(COUNTIF({rng},"NON OK")=0,"SEMELLE VÉRIFIÉE","SEMELLE NON VÉRIFIÉE")&" : A × B × H = "'
                 f'&TEXT({A},"0.00")&" × "&TEXT({B},"0.00")&" × "&TEXT({H},"0.00")&" m ; // A : "&{concl[0]}&" ; // B : "&{concl[1]}',
           YEL, size=12)
    S.ws.row_dimensions[r].height = 30
    S.ws.conditional_formatting.add(f"B4:K{r}", CellIsRule(operator="equal", formula=['"NON OK"'], fill=PINK))
    S.ws.conditional_formatting.add(f"E{r}", FormulaRule(formula=[f'COUNTIF({rng},"NON OK")>0'], fill=PINK))
    ws = S.ws
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def export_xlsx(prj, path, formulas=True):
    """One sheet per footing type (prj must have been run). formulas: print the formula row under each header."""
    wb = Workbook()
    wb.remove(wb.active)
    for res in prj.results:
        ws = wb.create_sheet(f"Semelle {res.gi.name}"[:31])
        footing_sheet(ws, prj, res, formulas)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path
