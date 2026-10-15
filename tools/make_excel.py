"""Build the clean calculation workbook « Semelle_isolee_RPA2024.xlsx » (formulas only, Excel 2007 functions).

Same method as the application (engine/), so both can be cross-checked:
DTR BC 2.331 (portance, diagramme trapézoïdal, règle des 3/4), RPA 2024 (§7.3.2, §7.4, §10.1.1, §10.1.4, §10.1.5),
BAEL 91 mod.99 / CBA 93 (bielles + moments S1 à 0,35 b, ELU / ELA / ELS, non-fragilité, poinçonnement).
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "Semelle_isolee_RPA2024.xlsx"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF", size=10)
BLACK = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F4E79")
H2 = Font(name=FONT, size=11, bold=True, color="FFFFFF")
GREY = Font(name=FONT, size=9, italic=True, color="666666")
INFILL = PatternFill("solid", fgColor="FFF2CC")
HFILL = PatternFill("solid", fgColor="1F4E79")
SUBFILL = PatternFill("solid", fgColor="D9E2F3")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
NROWS = 60  # load rows

wb = Workbook()


def name(n, sheet, cell):
    ref = f"'{sheet}'!${''.join(ch for ch in cell if ch.isalpha())}${''.join(ch for ch in cell if ch.isdigit())}"
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def header(ws, row, text, ncol=4):
    ws.cell(row, 1, text).font = H2
    for c in range(1, ncol + 1):
        ws.cell(row, c).fill = HFILL


def setw(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ====================================================================== Tables RPA 2024
tb = wb.active
tb.title = "Tables_RPA2024"
setw(tb, [22, 12, 12, 12, 46])
tb["A1"] = "Tables RPA 2024"
tb["A1"].font = TITLE
tb["A2"] = ("Valeurs du RPA 2024 (DTR B-C 2.48) : A tableau 3.3 ; S tableaux 3.5 (zones I-III) et 3.4 (zones IV-VI) ; "
            "I tableau 3.11 ; ψ tableau 4.2 ; α éq. (10.2).")
tb["A2"].font = GREY
header(tb, 4, "Zone → A (tab. 3.3) ; ρmin poteaux / espacement max (§7.4)", 5)
for j, h in enumerate(["Zone", "A", "ρmin fût (%)", "s max fût (m)"], 1):
    tb.cell(5, j, h).font = BOLD
zones = [("0", 0.00, 0.8, 0.20), ("I", 0.07, 0.8, 0.20), ("II", 0.10, 0.8, 0.20), ("III", 0.15, 0.9, 0.20),
         ("IV", 0.20, 1.0, 0.15), ("V", 0.25, 1.0, 0.15), ("VI", 0.30, 1.0, 0.15)]
for i, z in enumerate(zones, 6):
    for j, v in enumerate(z, 1):
        tb.cell(i, j, v).font = BLACK
header(tb, 14, "Groupe d'importance → I (tab. 3.11)", 5)
for i, (g, v) in enumerate([("1A", 1.40), ("1B", 1.20), ("2", 1.00), ("3", 0.80)], 15):
    tb.cell(i, 1, g).font = BLACK
    tb.cell(i, 2, v).font = BLACK
header(tb, 20, "Site → S (tab. 3.5 / 3.4) ; α longrines (éq. 10.2) ; section min longrine", 5)
for j, h in enumerate(["Site", "S zones I-III (t. 3.5)", "S zones IV-VI (t. 3.4)", "α", "b × h min (m)"], 1):
    tb.cell(21, j, h).font = BOLD
for i, row in enumerate([("S1", 1.00, 1.00, 0, "0,25 × 0,30"), ("S2", 1.30, 1.20, 0.3, "0,25 × 0,30"),
                         ("S3", 1.55, 1.30, 0.4, "0,25 × 0,30"), ("S4", 1.80, 1.35, 0.6, "0,30 × 0,30")], 22):
    for j, v in enumerate(row, 1):
        tb.cell(i, j, v).font = BLACK
header(tb, 27, "Sections d'acier (cm²)", 5)
for i, (p, a) in enumerate([(8, 0.503), (10, 0.785), (12, 1.131), (14, 1.539), (16, 2.011), (20, 3.142), (25, 4.909)], 28):
    tb.cell(i, 1, p).font = BLACK
    tb.cell(i, 2, a).font = BLACK

# ====================================================================== Données
d = wb.create_sheet("Données", 0)
setw(d, [46, 14, 10, 70])
d["A1"] = "SEMELLE ISOLÉE – DONNÉES"
d["A1"].font = TITLE
d["A2"] = "Cellules en bleu sur fond jaune = données à saisir. Tout le reste est calculé."
d["A2"].font = GREY
row = 4


def sec(title):
    global row
    row += 1
    header(d, row, title)
    row += 1


def inp(label, nm, val, unit="", src="", fmt="0.00", formula=False):
    global row
    d.cell(row, 1, label).font = BLACK
    c = d.cell(row, 2, val)
    c.number_format = fmt
    c.border = BOX
    if formula:
        c.font = BLACK
    else:
        c.font = BLUE
        c.fill = INFILL
    d.cell(row, 3, unit).font = BLACK
    d.cell(row, 4, src).font = GREY
    name(nm, "Données", f"B{row}")
    row += 1
    return f"B{row - 1}"


sec("Matériaux (BAEL 91 mod.99 / CBA 93 / RPA 2024)")
inp("fc28", "f_c28", 25, "MPa", "RPA 2024 §7.2 : béton ≥ C20/25 hors zone 0", "0")
inp("fe", "fe", 500, "MPa", "RPA 2024 §7.2 : HA 400 ≤ fe ≤ 600", "0")
inp("Fissuration : 1 = FPP, 2 = FP, 3 = FTP", "Fiss", 2, "", "", "0")
inp("Enrobage", "enr", 0.05, "m", "", "0.000")
inp("γ béton", "gc", 25, "kN/m³", "", "0.0")
inp("γb durable", "gb", 1.5, "", "BAEL A.4.3.41", "0.00")
inp("γb situation sismique", "gb_acc", 1.2, "", "RPA 2024 §7.3.2", "0.00")
inp("γs durable", "gs", 1.15, "", "BAEL A.4.3.2", "0.00")
inp("γs situation sismique", "gs_acc", 1.0, "", "RPA 2024 §7.3.2", "0.00")
inp("η (HA)", "eta", 1.6, "", "", "0.0")
inp("ft28 = 0,6 + 0,06·fc28", "f_t28", "=0.6+0.06*f_c28", "MPa", "BAEL A.2.1.12", "0.00", True)
inp("fbu = 0,85·fc28/γb", "fbu", "=0.85*f_c28/gb", "MPa", "BAEL A.4.3.41", "0.00", True)
inp("fbu sismique = 0,85·fc28/γb,acc", "fbu_acc", "=0.85*f_c28/gb_acc", "MPa", "RPA 2024 §7.3.2", "0.00", True)
inp("σs = fe/γs", "ss", "=fe/gs", "MPa", "", "0.0", True)
inp("σs sismique = fe/γs,acc", "ss_acc", "=fe/gs_acc", "MPa", "", "0.0", True)
inp("σ̄s ELS (0 si FPP : pas de limite)", "ssbar",
    "=IF(Fiss=1,0,IF(Fiss=2,1,0.8)*MIN(2/3*fe,MAX(0.5*fe,110*SQRT(eta*f_t28))))", "MPa",
    "BAEL 91 mod.99 A.4.5.33 / A.4.5.34", "0.0", True)
inp("τs = 0,6·ψs²·ft28 (ψs = 1,5)", "tau_s", "=0.6*1.5^2*f_t28", "MPa", "BAEL A.6.1.21", "0.00", True)

sec("Sol (DTR BC 2.331 / RPA 2024 §10.1.4-10.1.5)")
inp("σadm (rapport de sol)", "sadm", 2.0, "bar", "", "0.00")
inp("Profondeur d'ancrage D (TN → sous-face)", "Dan", 1.5, "m", "", "0.00")
inp("γ remblai", "gr", 20, "kN/m³", "", "0.0")
inp("k ELU : q_ELU = k·σadm", "k_elu", 1.5, "", "Fs 3 → 2 sur ql", "0.00")
inp("k sismique : q_ACC = k·σadm", "k_acc", 1.5, "", "RPA 2024 §10.1.4 : Fs = 2,0 sur Ql → 1,5·σadm", "0.00")
inp("φ' à l'interface (glissement)", "phis", 30, "°", "DTR BC 2.33.1", "0.0")
inp("q_ELS", "q_els", "=sadm*100", "kPa", "", "0.0", True)
inp("q_ELU", "q_elu", "=k_elu*q_els", "kPa", "", "0.0", True)
inp("q_ACC", "q_acc", "=k_acc*q_els", "kPa", "", "0.0", True)

sec("Géométrie")
inp("Fût ap (selon x)", "aP", 0.5, "m", "", "0.00")
inp("Fût bp (selon y)", "bP", 0.5, "m", "", "0.00")
inp("h0 : arase du fût au-dessus du TN", "h0", 0.15, "m", "", "0.00")
inp("A semelle (selon x)", "Asem", 2.15, "m", "voir A conseillé plus bas", "0.00")
inp("B semelle (selon y)", "Bsem", 2.15, "m", "", "0.00")
inp("h semelle", "hsem", 0.50, "m", "voir h minimal plus bas", "0.00")
inp("φ supposé pour d", "phid", 0.014, "m", "", "0.000")
inp("Espacement imposé des barres", "s_imp", 0.15, "m", "", "0.00")
inp("φ minimal semelle", "phimin", 12, "mm", "", "0")
inp("Hauteur du fût Hp = D − h + h0", "Hp", "=Dan-hsem+h0", "m", "", "0.00", True)
inp("Bras de levier des efforts horizontaux H = D + h0", "Hlev", "=Dan+h0", "m", "My,b = My + Vx·H ; Mx,b = Mx − Vy·H", "0.00", True)
inp("Poids semelle Wf = γc·A·B·h", "Wf", "=gc*Asem*Bsem*hsem", "kN", "", "0.00", True)
inp("Poids fût Wp = γc·ap·bp·Hp", "Wp", "=gc*aP*bP*Hp", "kN", "", "0.00", True)
inp("Poids remblai Ws = γr·(A·B − ap·bp)·(D − h)", "Ws", "=gr*(Asem*Bsem-aP*bP)*MAX(Dan-hsem,0)", "kN",
    "Erreur corrigée : le fichier d'origine prenait D au lieu de D − h", "0.00", True)
inp("Poids total W", "Wtot", "=Wf+Wp+Ws", "kN", "", "0.00", True)
inp("Charge sur console w = γc·h + γr·(D − h)", "wq", "=gc*hsem+gr*MAX(Dan-hsem,0)", "kPa", "", "0.00", True)
inp("dx (barres // A)", "dX", "=IF(Asem>=Bsem,hsem-enr-phid/2,hsem-enr-1.5*phid)", "m", "lit inférieur = plus grand côté", "0.000", True)
inp("dy (barres // B)", "dY", "=IF(Asem>=Bsem,hsem-enr-1.5*phid,hsem-enr-phid/2)", "m", "", "0.000", True)
inp("l1x = A/2 − 0,35·ap (section S1)", "l1_X", "=Asem/2-0.35*aP", "m", "DTU 13.12", "0.000", True)
inp("l1y = B/2 − 0,35·bp", "l1_Y", "=Bsem/2-0.35*bP", "m", "", "0.000", True)
inp("d requis (rigidité) = max((A − ap)/4 ; (B − bp)/4)", "d_req", "=MAX((Asem-aP)/4,(Bsem-bP)/4)", "m", "DTU 13.12", "0.000", True)
inp("h minimal = d requis + c + 1,5·φ (arrondi 5 cm)", "h_min", "=CEILING(d_req+enr+1.5*phid,0.05)", "m", "", "0.00", True)
inp("A conseillé – portance seule (homothétique) = √(1,15·Nser·ap/(bp·(q_ELS − 20·D)))", "A_cons",
    "=CEILING(SQRT(1.15*MAX(Vérifications!$Z$4,1)*aP/(bP*MAX(q_els-20*Dan,10))),0.05)", "m",
    "pré-dimensionnement ; vérifier ensuite toutes les conditions", "0.00", True)

sec("Coefficients de stabilité / poids propres")
inp("γG défavorable ELU (portance)", "gu_elu", 1.35, "", "", "0.00")
inp("γG favorable ELU (stabilité)", "gf_elu", 1.0, "", "", "0.00")
inp("γG ELS", "g_els", 1.0, "", "", "0.00")
inp("γG défavorable sismique", "gu_acc", 1.0, "", "", "0.00")
inp("γG favorable sismique", "gf_acc", 1.0, "", "RPA 2024 §4.5.1 (poids total)", "0.00")
inp("Fs renversement (ELS / ELU)", "Fs_ovt", 1.5, "", "pratique", "0.00")
inp("Fs renversement sismique", "Fs_ovt_acc", 1.3, "", "RPA 2024 §5.5", "0.00")
inp("Fs soulèvement ELS", "Fs_upl_els", 1.5, "", "", "0.00")
inp("Fs soulèvement ELU / sismique", "Fs_upl", 1.0, "", "", "0.00")
inp("Fs glissement ELS / ELU", "Fs_sld", 1.5, "", "DTR BC 2.33.1", "0.00")
inp("Fs glissement sismique", "Fs_sld_acc", 1.25, "", "RPA 2024 §5.5 / §10.1.5", "0.00")

sec("Séisme (RPA 2024)")
inp("Zone (0, I, II, III, IV, V, VI)", "zone", "II", "", "§3.1 / annexe A", "@")
inp("Site (S1…S4)", "site", "S2", "", "", "@")
inp("Groupe d'importance (1A, 1B, 2, 3)", "groupe", "2", "", "", "@")
inp("A", "Azone", "=INDEX(Tables_RPA2024!$B$6:$B$12,MATCH(zone,Tables_RPA2024!$A$6:$A$12,0))", "", "tab. 3.3", "0.00", True)
inp("I", "Iimp", "=INDEX(Tables_RPA2024!$B$15:$B$18,MATCH(groupe,Tables_RPA2024!$A$15:$A$18,0))", "", "tab. 3.11", "0.00", True)
inp("S", "Ssite", '=INDEX(IF(OR(zone="I",zone="II",zone="III"),Tables_RPA2024!$B$22:$B$25,Tables_RPA2024!$C$22:$C$25),'
    'MATCH(site,Tables_RPA2024!$A$22:$A$25,0))', "", "tab. 3.5 / 3.4", "0.00", True)
inp("R (coefficient de comportement)", "Rcoef", 4.5, "", "tab. 3.18", "0.00")
inp("Q_F (facteur de qualité)", "Qf", 1.25, "", "§3.8", "0.00")
inp("Ω = R/Q_F (fondations isolées)", "Omega", "=Rcoef/Qf", "", "éq. (5.6) ; 1,4 fondations communes éq. (5.7)", "0.00", True)
inp("ψ (Q en situation sismique)", "psi", 0.5, "", "tab. 4.2 : hangars 0,50 (habitation 0,30)", "0.00")
inp("ρmin fût selon zone", "rho_ped", "=INDEX(Tables_RPA2024!$C$6:$C$12,MATCH(zone,Tables_RPA2024!$A$6:$A$12,0))",
    "%", "RPA 2024 §7.4", "0.0", True)

for c, opts in (("Fiss", "1,2,3"),):
    pass
dv = DataValidation(type="list", formula1='"0,I,II,III,IV,V,VI"', allow_blank=False)
d.add_data_validation(dv)
dv.add(d[wb.defined_names["zone"].attr_text.split("!")[1].replace("$", "")])
dv2 = DataValidation(type="list", formula1='"S1,S2,S3,S4"')
d.add_data_validation(dv2)
dv2.add(d[wb.defined_names["site"].attr_text.split("!")[1].replace("$", "")])
dv3 = DataValidation(type="list", formula1='"1A,1B,2,3"')
d.add_data_validation(dv3)
dv3.add(d[wb.defined_names["groupe"].attr_text.split("!")[1].replace("$", "")])
dv4 = DataValidation(type="list", formula1='"1,2,3"')
d.add_data_validation(dv4)
dv4.add(d[wb.defined_names["Fiss"].attr_text.split("!")[1].replace("$", "")])
d.freeze_panes = "A4"

# ====================================================================== Combinaisons (helper)
cb = wb.create_sheet("Combinaisons", 1)
setw(cb, [34, 12, 12, 12, 12, 12, 40])
cb["A1"] = "GÉNÉRATEUR DE COMBINAISONS (aide) – à copier en valeurs dans « Efforts »"
cb["A1"].font = TITLE
cb["A2"] = ("Saisir les efforts des cas simples (réactions déjà converties : N > 0 compression). "
            "Copier ensuite le bloc résultat → Efforts (Collage spécial : valeurs).")
cb["A2"].font = GREY
cases = ["G", "Q", "Ex", "Ey"]
for j, h in enumerate(["Cas simple", "N (kN)", "Vx (kN)", "Vy (kN)", "Mx (kN·m)", "My (kN·m)"], 1):
    cb.cell(4, j, h).font = BOLD
    cb.cell(4, j).fill = SUBFILL
ex = {"G": (39.06 + 15.0, -7.09, 0.52, -0.78, -23.86), "Q": (8.0, -1.5, 0.1, -0.2, -6.0),
      "Ex": (1.13, -4.35, -0.02, -0.03, -19.31), "Ey": (0.8, -0.3, -3.1, -9.5, -0.5)}
for i, cs in enumerate(cases, 5):
    cb.cell(i, 1, cs).font = BOLD
    for j in range(5):
        c = cb.cell(i, 2 + j, ex[cs][j])
        c.font = BLUE
        c.fill = INFILL
        c.number_format = "0.00"
cb["G5"] = "Exemple – remplacer par vos valeurs"
cb["G5"].font = GREY
cb["A10"] = "Facteur sismique λ (Ω = R/Q_F isolées, 1,4 communes, 1,0 horizontale)"
cb["B10"] = "=Omega"
cb["B10"].font = BLUE
cb["B10"].fill = INFILL
cb["B10"].number_format = "0.00"
cb["C10"] = "ψ ="
cb["D10"] = "=psi"
for j, h in enumerate(["Combinaison", "Type", "N", "Vx", "Vy", "Mx", "My"], 1):
    cb.cell(12, j, h).font = BOLD
    cb.cell(12, j).fill = SUBFILL
combos = [("ELU : 1,35G + 1,5Q", "ELU", 1.35, 1.5, 0, 0), ("ELS : G + Q", "ELS", 1, 1, 0, 0)]
for s1 in (1, -1):
    for a, b in ((1, 0.3), (0.3, 1)):
        for s2 in (1, -1):
            nm = f"G+ψQ{'+' if s1 > 0 else '−'}λ({a:g}Ex{'+' if s2 > 0 else '−'}{b:g}Ey)"
            combos.append((nm, "ACC", 1, "psi", s1 * a, s1 * b * s2))
for i, (nm, t, fg, fq, fx, fy) in enumerate(combos, 13):
    cb.cell(i, 1, nm).font = BLACK
    cb.cell(i, 2, t).font = BLACK
    for j in range(5):
        col = get_column_letter(2 + j)
        fqs = "$D$10" if fq == "psi" else str(fq)
        if t == "ACC":
            f = f"={fg}*{col}$5+{fqs}*{col}$6+$B$10*({fx}*{col}$7+{fy}*{col}$8)"
        else:
            f = f"={fg}*{col}$5+{fqs}*{col}$6"
        c = cb.cell(i, 3 + j, f)
        c.number_format = "0.00"
        c.font = BLACK

# ====================================================================== Efforts (input table)
ef = wb.create_sheet("Efforts", 2)
setw(ef, [6, 40, 9, 12, 12, 12, 12, 12])
ef["A1"] = "EFFORTS EN PIED DE POTEAU (combinaisons) – N > 0 compression, kN / kN·m"
ef["A1"].font = TITLE
ef["A2"] = ("Une ligne par combinaison et par appui (enveloppe de tous les appuis de la semelle type). Type = ELU, ELS ou ACC. "
            "Robot : réactions → action sur la semelle : N = FZ ; Vx = −FX ; Vy = −FY ; Mx = −MX ; My = −MY.")
ef["A2"].font = GREY
for j, h in enumerate(["N°", "Combinaison / appui", "Type", "N", "Vx", "Vy", "Mx", "My"], 1):
    ef.cell(5, j, h).font = BOLD
    ef.cell(5, j).fill = SUBFILL
example = [("1864 – ELU", "ELU", 68.3, -13.9, 0.97, -1.46, -46.75),
           ("1864 – ELS", "ELS", 50.59, -10.29, 0.72, -1.08, -34.62),
           ("1864 – 29 Acc+λ(Ex+0,3Ey)", "ACC", 52.14, -16.37, 0.69, -1.11, -61.62),
           ("1864 – 30 Acc+λ(0,3Ex+Ey)", "ACC", 51.03, -12.11, 0.71, -1.08, -42.70),
           ("1864 – 33 Acc−λ(Ex+0,3Ey)", "ACC", 48.97, -4.19, 0.74, -1.03, -7.55),
           ("1864 – 34 Acc−λ(0,3Ex+Ey)", "ACC", 50.08, -8.46, 0.73, -1.06, -26.48)]
dvt = DataValidation(type="list", formula1='"ELU,ELS,ACC"')
ef.add_data_validation(dvt)
for i in range(NROWS):
    r = 6 + i
    ef.cell(r, 1, i + 1).font = GREY
    vals = example[i] if i < len(example) else ("", "", None, None, None, None, None)
    for j, v in enumerate(vals, 2):
        c = ef.cell(r, j, v)
        c.font = BLUE
        c.fill = INFILL
        c.border = BOX
        if j >= 4:
            c.number_format = "0.00"
    dvt.add(f"C{r}")
ef["J5"] = "Lignes d'exemple (appui 1864 du modèle Robot) : à remplacer par vos efforts."
ef["J5"].font = GREY
ef.freeze_panes = "C6"

# ====================================================================== Vérifications (geotechnical)
ve = wb.create_sheet("Vérifications", 3)
cols = ["Combinaison", "Type", "γG déf.", "γG fav.", "Nt déf. (kN)", "Nt fav. (kN)", "My,b (kN·m)", "Mx,b (kN·m)",
        "ex (m)", "ey (m)", "Diagr. x", "σ3/4 x (kPa)", "Diagr. y", "σ3/4 y (kPa)", "q lim (kPa)", "Taux portance",
        "Portance", "Taux renversement", "Renversement", "Taux soulèvement", "Soulèvement", "H (kN)",
        "R glissement (kN)", "Taux glissement", "Glissement", "N ELS (aide)"]
setw(ve, [34, 7] + [11] * (len(cols) - 2))
ve["A1"] = "VÉRIFICATIONS GÉOTECHNIQUES ET DE STABILITÉ (toutes les combinaisons)"
ve["A1"].font = TITLE
ve["A2"] = ("σ3/4 = (3σmax + σmin)/4 = N/(A·B)·(1 + 3e/A) si e ≤ A/6 ; = N/(2·B·(A/2 − e)) si A/6 < e < A/2 (triangle). "
            "Renversement : Fs = N·(A/2)/M ≥ 1,3 sous ACC (RPA 2024 §5.5), ≥ 1,5 en ELS/ELU. Glissement : Fs = N·tanφ'/H.")
ve["A2"].font = GREY
ve["A3"] = "MAX →"
ve["A3"].font = BOLD
for j, h in enumerate(cols, 1):
    c = ve.cell(5, j, h)
    c.font = BOLD
    c.fill = SUBFILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
ve.row_dimensions[5].height = 42
for i in range(NROWS):
    r, e = 6 + i, 6 + i
    E = f"Efforts!$D{e}"
    f = {
        "A": f'=IF({E}="","",Efforts!B{e})',
        "B": f'=IF({E}="","",Efforts!C{e})',
        "C": f'=IF(B{r}="","",IF(B{r}="ELU",gu_elu,IF(B{r}="ACC",gu_acc,g_els)))',
        "D": f'=IF(B{r}="","",IF(B{r}="ELU",gf_elu,IF(B{r}="ACC",gf_acc,g_els)))',
        "E": f'=IF(B{r}="","",{E}+C{r}*Wtot)',
        "F": f'=IF(B{r}="","",{E}+D{r}*Wtot)',
        "G": f'=IF(B{r}="","",Efforts!H{e}+Efforts!E{e}*Hlev)',
        "H": f'=IF(B{r}="","",Efforts!G{e}-Efforts!F{e}*Hlev)',
        "I": f'=IF(B{r}="","",IF(E{r}>0,ABS(G{r})/E{r},99))',
        "J": f'=IF(B{r}="","",IF(E{r}>0,ABS(H{r})/E{r},99))',
        "K": f'=IF(B{r}="","",IF(E{r}<=0,"N≤0",IF(I{r}<=Asem/6,"trapèze",IF(I{r}<Asem/2,"triangle","instable"))))',
        "L": f'=IF(B{r}="","",IF(K{r}="trapèze",E{r}/(Asem*Bsem)*(1+3*I{r}/Asem),IF(K{r}="triangle",E{r}/(2*Bsem*(Asem/2-I{r})),IF(K{r}="N≤0",0,1E9))))',
        "M": f'=IF(B{r}="","",IF(E{r}<=0,"N≤0",IF(J{r}<=Bsem/6,"trapèze",IF(J{r}<Bsem/2,"triangle","instable"))))',
        "N": f'=IF(B{r}="","",IF(M{r}="trapèze",E{r}/(Asem*Bsem)*(1+3*J{r}/Bsem),IF(M{r}="triangle",E{r}/(2*Asem*(Bsem/2-J{r})),IF(M{r}="N≤0",0,1E9))))',
        "O": f'=IF(B{r}="","",IF(B{r}="ELU",q_elu,IF(B{r}="ACC",q_acc,q_els)))',
        "P": f'=IF(B{r}="","",MAX(L{r},N{r})/O{r})',
        "Q": f'=IF(B{r}="","",IF(P{r}<=1,"OK","NON OK"))',
        "R": f'=IF(B{r}="","",IF(F{r}<=0,99,IF(B{r}="ACC",Fs_ovt_acc,Fs_ovt)*MAX(ABS(G{r})/(F{r}*Asem/2),ABS(H{r})/(F{r}*Bsem/2))))',
        "S": f'=IF(B{r}="","",IF(R{r}<=1,"OK","NON OK"))',
        "T": f'=IF(B{r}="","",IF({E}<0,IF(B{r}="ELS",Fs_upl_els,Fs_upl)*ABS({E})/(D{r}*Wtot),0))',
        "U": f'=IF(B{r}="","",IF(T{r}<=1,"OK","NON OK"))',
        "V": f'=IF(B{r}="","",SQRT(Efforts!E{e}^2+Efforts!F{e}^2))',
        "W": f'=IF(B{r}="","",MAX(F{r},0)*TAN(RADIANS(phis)))',
        "X": f'=IF(B{r}="","",IF(V{r}<0.000001,0,IF(W{r}<=0,99,IF(B{r}="ACC",Fs_sld_acc,Fs_sld)*V{r}/W{r})))',
        "Y": f'=IF(B{r}="","",IF(X{r}<=1,"OK","NON OK"))',
        "Z": f'=IF(B{r}="ELS",{E},"")',
    }
    for col, formula in f.items():
        c = ve[f"{col}{r}"]
        c.value = formula
        c.font = BLACK
        c.number_format = "0.00" if col not in ("P", "R", "T", "X") else "0%"
for col in ("E", "L", "N", "P", "R", "T", "V", "X"):
    c = ve[f"{col}3"]
    c.value = f"=MAX({col}6:{col}{5 + NROWS})"
    c.font = BOLD
    c.number_format = "0%" if col in ("P", "R", "T", "X") else "0.00"
ve["Z3"] = "N ELS max"
ve["Z4"] = f"=MAX(Z6:Z{5 + NROWS})"
red = PatternFill("solid", fgColor="F8CBAD")
green = PatternFill("solid", fgColor="C6EFCE")
for col in ("Q", "S", "U", "Y"):
    rng = f"{col}6:{col}{5 + NROWS}"
    ve.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NON OK"'], fill=red))
    ve.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=green))
ve.freeze_panes = "C6"

# ====================================================================== Ferraillage (per row, both directions)
fe_ = wb.create_sheet("Ferraillage", 4)
fe_["A1"] = "FERRAILLAGE DE LA SEMELLE – méthode des bielles et méthode des moments (section S1 à 0,35·b)"
fe_["A1"].font = TITLE
fe_["A2"] = ("Bielles (e0 ≤ L/6) : As = N'·(L − l)/(8·d·σs), N' = Nu·(1 + 3e0/L), Nu = N + γG·Wp. "
             "Moments : M1 = B·∫σ(s)(l1 − s)ds − w·B·l1²/2 (réaction du sol sous Nt, poids semelle + remblai déduit) ; "
             "ELU/ACC : μ, α, z ; ELS (FP/FTP) : As = M1/(0,9·d·σ̄s).")
fe_["A2"].font = GREY
common = ["Combinaison", "Type", "γG", "σs (MPa)", "fbu (MPa)", "Nu fût (kN)", "Nt (kN)"]
dircols = ["M base", "e0", "As bielles", "e (Nt)", "σmax", "k", "s fin", "M sol", "M1", "μ", "As moments",
           "As inf.", "As sup."]
allcols = common + [f"X: {c}" for c in dircols] + [f"Y: {c}" for c in dircols] + ["P'u (kN)", "P lim (kN)", "Taux poinç."]
setw(fe_, [34, 7] + [10] * (len(allcols) - 2))
for j, h in enumerate(allcols, 1):
    c = fe_.cell(5, j, h)
    c.font = BOLD
    c.fill = SUBFILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
fe_.row_dimensions[5].height = 42


def L(j):
    return get_column_letter(j)


for i in range(NROWS):
    r = 6 + i
    e = 6 + i
    N = f"Efforts!$D{e}"
    base = {
        1: f'=IF({N}="","",Efforts!B{e})',
        2: f'=IF({N}="","",Efforts!C{e})',
        3: f'=IF(B{r}="","",IF(B{r}="ELU",gu_elu,IF(B{r}="ACC",gu_acc,g_els)))',
        4: f'=IF(B{r}="","",IF(B{r}="ELS",IF(Fiss=1,0,ssbar),IF(B{r}="ACC",ss_acc,ss)))',
        5: f'=IF(B{r}="","",IF(B{r}="ACC",fbu_acc,fbu))',
        6: f'=IF(B{r}="","",{N}+C{r}*Wp)',
        7: f'=IF(B{r}="","",{N}+C{r}*Wtot)',
    }
    for j, f in base.items():
        fe_.cell(r, j, f).font = BLACK
    for k, (Ld, Wd, ld, l1, dd, mom) in enumerate((("Asem", "Bsem", "aP", "l1_X", "dX", f"Efforts!H{e}+Efforts!E{e}*Hlev"),
                                                    ("Bsem", "Asem", "bP", "l1_Y", "dY", f"Efforts!G{e}-Efforts!F{e}*Hlev"))):
        c0 = 8 + k * len(dircols)
        cM, ce0, cbi, cet, csm, ck, cse, cms, cm1, cmu, caf, cinf, csup = [L(c0 + t) for t in range(len(dircols))]
        fs = {
            cM: f'=IF(B{r}="","",{mom})',
            ce0: f'=IF(B{r}="","",IF(F{r}>0,ABS({cM}{r})/F{r},99))',
            cbi: f'=IF(OR(B{r}="",D{r}=0),"",IF(AND(F{r}>0,{ce0}{r}<={Ld}/6),10*F{r}*(1+3*{ce0}{r}/{Ld})*({Ld}-{ld})/(8*{dd}*D{r}),""))',
            cet: f'=IF(B{r}="","",IF(G{r}>0,ABS({cM}{r})/G{r},-1))',
            csm: f'=IF(B{r}="","",IF({cet}{r}<0,0,IF({cet}{r}<={Ld}/6,G{r}/({Ld}*{Wd})*(1+6*{cet}{r}/{Ld}),IF({cet}{r}<{Ld}/2,2*G{r}/({Wd}*3*({Ld}/2-{cet}{r})),1E9))))',
            ck: f'=IF(B{r}="","",IF({cet}{r}<0,0,IF({cet}{r}<={Ld}/6,({csm}{r}-G{r}/({Ld}*{Wd})*(1-6*{cet}{r}/{Ld}))/{Ld},IF({cet}{r}<{Ld}/2,{csm}{r}/(3*({Ld}/2-{cet}{r})),0))))',
            cse: f'=IF(B{r}="","",IF({ck}{r}<=0,{l1},MIN({l1},{csm}{r}/{ck}{r})))',
            cms: f'=IF(B{r}="","",IF({cet}{r}<0,0,{Wd}*({csm}{r}*({l1}*{cse}{r}-{cse}{r}^2/2)-{ck}{r}*({l1}*{cse}{r}^2/2-{cse}{r}^3/3))))',
            cm1: f'=IF(B{r}="","",{cms}{r}-C{r}*wq*{Wd}*{l1}^2/2)',
            cmu: f'=IF(OR(B{r}="",B{r}="ELS"),"",ABS({cm1}{r})/({Wd}*{dd}^2*E{r}*1000))',
            caf: f'=IF(OR(B{r}="",D{r}=0),"",IF(B{r}="ELS",10*ABS({cm1}{r})/(0.9*{dd}*D{r}),IF({cmu}{r}>=0.5,1E9,'
                 f'10*ABS({cm1}{r})/({dd}*(1-0.4*1.25*(1-SQRT(1-2*{cmu}{r})))*D{r}))))',
            cinf: f'=IF({caf}{r}="","",IF({cm1}{r}>=0,{caf}{r},0))',
            csup: f'=IF({caf}{r}="","",IF({cm1}{r}<0,{caf}{r},0))',
        }
        for col, f in fs.items():
            c = fe_[f"{col}{r}"]
            c.value = f
            c.font = BLACK
            c.number_format = "0.000" if col in (ce0, cet, cse, cmu) else "0.00"
    cp = 8 + 2 * len(dircols)
    fe_.cell(r, cp, f'=IF(OR(B{r}="",B{r}="ELS"),"",IF(F{r}<=0,0,F{r}*(1-MIN(aP+2*hsem,Asem)*MIN(bP+2*hsem,Bsem)/(Asem*Bsem))))')
    fe_.cell(r, cp + 1, f'=IF(OR(B{r}="",B{r}="ELS"),"",0.045*2*(aP+bP+2*hsem)*hsem*f_c28/IF(B{r}="ACC",gb_acc,gb)*1000)')
    fe_.cell(r, cp + 2, f'=IF({L(cp)}{r}="","",{L(cp)}{r}/{L(cp + 1)}{r})')
    for t in range(3):
        fe_.cell(r, cp + t).font = BLACK
        fe_.cell(r, cp + t).number_format = "0.00" if t < 2 else "0%"
fe_.freeze_panes = "C6"
FX0, FY0 = 8, 8 + len(dircols)
colX = {n: L(FX0 + t) for t, n in enumerate(dircols)}
colY = {n: L(FY0 + t) for t, n in enumerate(dircols)}
CP = 8 + 2 * len(dircols)

# ====================================================================== Résultats
rs = wb.create_sheet("Résultats", 5)
setw(rs, [52, 14, 14, 14, 14, 40])
rs["A1"] = "RÉSULTATS ET FERRAILLAGE RETENU"
rs["A1"].font = TITLE
R = 3
rng = lambda col: f"Ferraillage!{col}6:{col}{5 + NROWS}"  # noqa: E731


def line(label, fx, fy="", unit="", fmt="0.00", bold=False, note=""):
    global R
    rs.cell(R, 1, label).font = BOLD if bold else BLACK
    for j, f in ((2, fx), (3, fy)):
        if f != "":
            c = rs.cell(R, j, f)
            c.font = BOLD if bold else BLACK
            c.number_format = fmt
    rs.cell(R, 4, unit).font = BLACK
    rs.cell(R, 6, note).font = GREY
    R += 1
    return R - 1


header(rs, R, "Vérifications (taux maximal sur toutes les combinaisons)", 6)
R += 1
rs.cell(R, 2, "Taux").font = BOLD
rs.cell(R, 3, "Statut").font = BOLD
R += 1
for lab, col in (("Portance du sol (DTR BC 2.331)", "P"), ("Renversement (RPA / Fs)", "R"),
                 ("Soulèvement", "T"), ("Glissement (RPA 2024 §10.1.5)", "X")):
    rr = line(lab, f"=Vérifications!{col}3", f'=IF(Vérifications!{col}3<=1,"OK","NON OK")', "", "0%")
rr = line("Rigidité : h ≥ h minimal", "=hsem", '=IF(hsem>=h_min-0.0001,"OK","NON OK")', "m",
          note="h minimal calculé dans Données")
rr = line("Poinçonnement (BAEL A.5.2.4)", f"=MAX({rng(L(CP + 2))})", f'=IF(B{R}<=1,"OK","NON OK")', "", "0%")
rs.cell(rr, 3).value = f'=IF(B{rr}<=1,"OK","NON OK")'
R += 1
header(rs, R, "Ferraillage de la semelle", 6)
R += 1
rs.cell(R, 2, "// A (x)").font = BOLD
rs.cell(R, 3, "// B (y)").font = BOLD
R += 1
r_bi = line("As bielles max (ELU / ACC / ELS)", f"=MAX({rng(colX['As bielles'])})", f"=MAX({rng(colY['As bielles'])})", "cm²")
r_mo = line("As moments max (nappe inférieure)", f"=MAX({rng(colX['As inf.'])})", f"=MAX({rng(colY['As inf.'])})", "cm²")
r_mi = line("As min non-fragilité = 0,23·b·d·ft28/fe", "=0.23*Bsem*dX*f_t28/fe*10000", "=0.23*Asem*dY*f_t28/fe*10000",
            "cm²", note="BAEL A.4.2.1")
r_as = line("As requis = max(bielles ; moments ; min)", f"=MAX(B{r_bi},B{r_mo},B{r_mi})", f"=MAX(C{r_bi},C{r_mo},C{r_mi})",
            "cm²", bold=True)
r_top = line("As nappe supérieure (soulèvement)", f"=IF(MAX({rng(colX['As sup.'])})>0,MAX(MAX({rng(colX['As sup.'])}),B{r_mi}),0)",
             f"=IF(MAX({rng(colY['As sup.'])})>0,MAX(MAX({rng(colY['As sup.'])}),C{r_mi}),0)", "cm²")
r_n = line("Nombre de barres (espacement imposé)", "=CEILING((Bsem-2*enr)/s_imp,1)+1", "=CEILING((Asem-2*enr)/s_imp,1)+1", "", "0")
r_sm = line("Espacement maximal (fissuration)", "=IF(Fiss=3,MIN(1.5*hsem,0.2),IF(Fiss=2,MIN(2*hsem,0.25),MIN(3*hsem,0.33)))",
            f"=B{R}", "m", note="BAEL A.4.5.33-34 / A.8.2.42")
rs.cell(r_sm, 3).value = f"=B{r_sm}"
# diameter choice table
R += 1
rs.cell(R, 1, "Choix du diamètre (premier φ ≥ φmin qui satisfait As)").font = BOLD
R += 1
t0 = R
for k, phi in enumerate((10, 12, 14, 16, 20, 25)):
    rr = R + k
    rs.cell(rr, 1, f"HA{phi} : section fournie (cm²) / convient (1 = oui)").font = BLACK
    a = f"INDEX(Tables_RPA2024!$B$28:$B$34,MATCH({phi},Tables_RPA2024!$A$28:$A$34,0))"
    rs.cell(rr, 2, f"=B${r_n}*{a}").number_format = "0.00"
    rs.cell(rr, 3, f"=C${r_n}*{a}").number_format = "0.00"
    rs.cell(rr, 4, f"=IF(AND({phi}>=phimin,B{rr}>=B${r_as}),1,0)")
    rs.cell(rr, 5, f"=IF(AND({phi}>=phimin,C{rr}>=C${r_as}),1,0)")
    rs.cell(rr, 6, phi)
R += 6
rs.cell(t0 - 1, 4, "ok // A").font = BOLD
rs.cell(t0 - 1, 5, "ok // B").font = BOLD
rs.cell(t0 - 1, 6, "φ").font = BOLD
r_phi = line("Diamètre retenu (mm)", f'=IFERROR(INDEX(F{t0}:F{t0 + 5},MATCH(1,D{t0}:D{t0 + 5},0)),"> HA25 : réduire s")',
             f'=IFERROR(INDEX(F{t0}:F{t0 + 5},MATCH(1,E{t0}:E{t0 + 5},0)),"> HA25 : réduire s")', "mm", "0", bold=True)
r_txt = line("Ferraillage retenu", f'=B{r_n}&" HA"&B{r_phi}&" e = "&TEXT(s_imp*100,"0")&" cm"',
             f'=C{r_n}&" HA"&C{r_phi}&" e = "&TEXT(s_imp*100,"0")&" cm"', "", "@", bold=True)
line("Espacement ≤ espacement maximal", f'=IF(s_imp<=B{r_sm}+0.0001,"OK","NON OK")', f'=IF(s_imp<=C{r_sm}+0.0001,"OK","NON OK")')
r_ls = line("Longueur de scellement ls = φ·fe/(4·τs)", f"=B{r_phi}/1000*fe/(4*tau_s)", f"=C{r_phi}/1000*fe/(4*tau_s)", "m",
            "0.000", note="BAEL A.6.1.22")
line("Disposition des ancrages", f'=IF(B{r_ls}>Asem/4,"crochets",IF(B{r_ls}>Asem/8,"barres droites","droites, arrêt 1 sur 2 possible"))',
     f'=IF(C{r_ls}>Bsem/4,"crochets",IF(C{r_ls}>Bsem/8,"barres droites","droites, arrêt 1 sur 2 possible"))',
     note="DTU 13.12 : ls > L/4 crochets ; L/8 < ls ≤ L/4 droites")
R += 1
header(rs, R, "Fût (amorce de poteau) – minimum RPA 2024 §7.4", 6)
R += 1
r_pm = line("As min fût = ρmin·ap·bp", "=rho_ped/100*aP*bP*10000", "", "cm²")
line("Barres (HA12 mini, ≥ 1 barre intermédiaire par face)", f'=MAX(8,2*CEILING(B{r_pm}/1.131/2,1))&" HA12 (ou équivalent)"', "", "")
R += 1
header(rs, R, "Dimensions", 6)
R += 1
line("A × B × h retenus", '=TEXT(Asem,"0.00")&" × "&TEXT(Bsem,"0.00")&" × "&TEXT(hsem,"0.00")&" m"', "", "", "@", bold=True)
line("A conseillé (pré-dimensionnement)", "=A_cons", "", "m")
line("h minimal (rigidité)", "=h_min", "", "m")
for col in ("C",):
    rgn = f"{col}3:{col}{R}"
    rs.conditional_formatting.add(rgn, CellIsRule(operator="equal", formula=['"NON OK"'], fill=red))
    rs.conditional_formatting.add(rgn, CellIsRule(operator="equal", formula=['"OK"'], fill=green))

# ====================================================================== Longrine
lg = wb.create_sheet("Longrine", 6)
setw(lg, [56, 14, 10, 50])
lg["A1"] = "LONGRINE – RPA 2024 §10.1.1"
lg["A1"].font = TITLE
row = 3
li = []


def lin(label, val, unit="", src="", fmt="0.00", inputcell=False):
    global row
    lg.cell(row, 1, label).font = BLACK
    c = lg.cell(row, 2, val)
    c.number_format = fmt
    c.border = BOX
    c.font = BLUE if inputcell else BLACK
    if inputcell:
        c.fill = INFILL
    lg.cell(row, 3, unit).font = BLACK
    lg.cell(row, 4, src).font = GREY
    row += 1
    return f"B{row - 1}"


b_ = lin("b", 0.30, "m", "", "0.00", True)
h_ = lin("h", 0.40, "m", "", "0.00", True)
L_ = lin("Portée libre L (entre nus des fûts)", 4.60, "m", "entre axes − ½ fût de chaque côté", "0.00", True)
N1 = lin("N appui 1 (situation sismique)", 52.14, "kN", "exemple – max ACC de l'appui", "0.00", True)
N2 = lin("N appui 2 (situation sismique)", 50.00, "kN", "exemple", "0.00", True)
qa = lin("Charge additionnelle (mur…)", 0, "kN/m", "", "0.00", True)
al = lin("α (site)", "=INDEX(Tables_RPA2024!$D$22:$D$25,MATCH(site,Tables_RPA2024!$A$22:$A$25,0))", "", "éq. (10.2)")
Nm = lin("N = (N1 + N2)/2", f"=({N1}+{N2})/2", "kN", "valeur moyenne des appuis solidarisés")
F_ = lin("F = α·(A·I·S)·N ≥ 20 kN", f"=MAX({al}*Azone*Iimp*Ssite*{Nm},20)", "kN", "éq. (10.1)")
Ast = lin("As,F = F/σs,acc", f"=10*{F_}/ss_acc", "cm²", "RPA 2024 §7.3.2 : γs = 1")
dd = lin("d = h − c − 1 cm", f"={h_}-enr-0.01", "m")
Mf = lin("M = 1,35·(γc·b·h + q)·L²/8", f"=1.35*(gc*{b_}*{h_}+{qa})*{L_}^2/8", "kN·m")
Asf = lin("As,f = M/(0,9·d·σs)", f"=10*{Mf}/(0.9*{dd}*ss)", "cm²")
Amin = lin("As,min = 0,6 %·b·h", f"=0.006*{b_}*{h_}*10000", "cm²", "RPA 2024 §10.1.1")
Asr = lin("As retenu = max(As,F + As,f ; As,min)", f"=MAX({Ast}+{Asf},{Amin})", "cm²")
row += 1
lg.cell(row, 1, "Choix des barres").font = BOLD
row += 1
opts = [(4, 12), (6, 12), (8, 12), (4, 14), (6, 14), (8, 14), (4, 16), (6, 16), (8, 16), (4, 20), (6, 20), (8, 20)]
r0 = row
for n, phi in opts:
    lg.cell(row, 1, f"{n} HA{phi}")
    lg.cell(row, 2, f"={n}*INDEX(Tables_RPA2024!$B$28:$B$34,MATCH({phi},Tables_RPA2024!$A$28:$A$34,0))").number_format = "0.00"
    lg.cell(row, 3, f"=IF(B{row}>={Asr},1,0)")
    lg.cell(row, 4, phi)
    row += 1
lin("Ferraillage retenu", f'=IFERROR(INDEX(A{r0}:A{row - 1},MATCH(1,C{r0}:C{row - 1},0)),"augmenter la section")', "", "", "@")
lin("Section ≥ minimum du site", f'=IF(AND({b_}>=IF(site="S4",0.3,0.25),{h_}>=0.3),"OK","NON OK")', "", "25×30 (S2,S3) / 30×30 (S4)", "@")
lin("Cadres : St < min(20 cm ; 15·φl)", f'="HA8, St = "&TEXT(MIN(20,15*IFERROR(INDEX(D{r0}:D{row - 2},MATCH(1,C{r0}:C{row - 2},0)),12)/10),"0")&" cm"', "", "", "@")
lin("Longrines exigées ?", '=IF(OR(site="S1",AND(site="S2",OR(zone="0",zone="I"))),"non exigées (§10.1.1 c)","exigées")', "", "", "@")

# ====================================================================== Audit of the original file
au = wb.create_sheet("Audit fichier d'origine", 7)
setw(au, [6, 30, 70, 60])
au["A1"] = "AUDIT DU FICHIER « Semelle-isolée RPA2024 (2) justeeeeeeeeeeeee.xlsx »"
au["A1"].font = TITLE
for j, h in enumerate(["N°", "Cellule(s)", "Problème", "Correction dans ce classeur"], 1):
    au.cell(3, j, h).font = BOLD
    au.cell(3, j).fill = SUBFILL
issues = [
    ("E142:E271 (σ sens X)", "Le cas triangulaire utilise $C$142 (N de l'ELU) au lieu du N de la ligne : toutes les combinaisons "
     "ELS / ACC sont vérifiées avec l'effort ELU.", "Vérifications!L : N de chaque ligne."),
    ("F9:G138 (excentricités)", "e = M/N calculé avec N du poteau seul, sans le poids de la semelle et des terres → e surestimé, "
     "diagramme faux ; e > A/2 donne σ < 0 : le fichier affiche « traction non vérifiée » pour l'ELU et l'ELS.",
     "e = M/Nt avec Nt = N + γG·W."),
    ("C9:E138 (efforts)", "Les efforts horizontaux Vx, Vy ne sont pas pris : le moment à la base de la semelle (M + V·H) est "
     "ignoré. Pour un portique encastré c'est la principale source de moment.", "My,b = My + Vx·H ; Mx,b = Mx − Vy·H."),
    ("D352:D361 (= I9…)", "My du calcul des armatures pointe vers la colonne I vide → My = 0 : la méthode des bielles est "
     "calculée sans excentricité (armatures sous-estimées).", "Ferraillage : moment réel de chaque combinaison."),
    ("B358:D361, C305:E308", "#REF! (références supprimées) pour ACC5 à ACC8 : poinçonnement et contraintes non calculés.",
     "Toutes les lignes calculées, sans référence brisée."),
    ("F303", "As // Y = MAX(F291:F300) mélange les deux directions (même valeur que // X).", "Directions X et Y séparées."),
    ("F291:F300, D311", "Armatures et poinçonnement ne prennent que ELU + ACC1…ACC4 (4 combinaisons sismiques sur 128).",
     "Toutes les combinaisons."),
    ("D291 (bielles)", "N' calculé avec N total (poids de la semelle et des terres inclus) : ce poids ne fléchit pas la semelle. "
     "Pas de contrôle e0 ≤ B/6 (domaine de validité des bielles).", "Nu = N + γG·Wp ; bielles seulement si e0 ≤ L/6, "
     "sinon méthode des moments."),
    ("H3 (poids des terres)", "γ·D·(A·B − a·b) : prend toute la profondeur D, la hauteur de la semelle est comptée deux fois "
     "(terre + béton) → poids stabilisant surestimé (non sécuritaire au renversement / soulèvement).", "Ws = γr·(A·B − ap·bp)·(D − h)."),
    ("Z12", "Poids du fût (amorce) non pris en compte (= 0).", "Wp = γc·ap·bp·Hp."),
    ("D5:F5 / D4:F4", "Dimensions saisies en dur (1 × 1 × 0,30) sans lien avec le pré-dimensionnement ; h proposé = 0,05 m.",
     "Données : A conseillé et h minimal calculés, contrôlés dans Résultats."),
    ("E291 (d)", "d = 0,9·h (au lieu de h − c − φ/2).", "dX, dY réels."),
    ("—", "Aucune vérification du renversement (RPA 2024 §5.5, Fs ≥ 1,3), du soulèvement ni du glissement (§10.1.5).",
     "Vérifications!R, T, X."),
    ("—", "Pas de condition de non-fragilité ni de contrôle de l'espacement / ancrage.", "Résultats."),
    ("L388:L395 (semelles excentrées)", "fc28 = 20 alors que la note est en 25 ; fbu(ELA) = fc28·0,85/(1,15·0,85) ; "
     "efforts N1 = 400, N2 = 600 saisis en dur (bloc d'exemple non relié).", "fbu,acc = 0,85·fc28/1,20 (RPA 2024)."),
    ("RECAP!B6:E9", "d = (B − 0,2)/4 et hmin avec 0,2 / 0,35 en dur (fût de 20 / 35 cm) au lieu des dimensions réelles ; "
     "C4:E5 = 1,7 en dur ; #REF! en F95:G95.", "Données / Résultats reliés."),
    ("Feuilles B1, A1, A2", "#REF! en B8 et C41:E44 (mêmes erreurs que la feuille principale).", "—"),
    ("Points corrects", "Diagramme trapézoïdal σ3/4 = N/(A·B)(1 + 3e/A), triangle N/(2B(A/2 − e)) ; limite sismique 1,5·σadm ; "
     "poinçonnement 0,045·uc·h·fc28/γb avec γb = 1,2 en sismique ; σ̄s ELS (CBA).", "Conservés."),
]
for i, (cell, prob, fix) in enumerate(issues, 4):
    au.cell(i, 1, i - 3).font = BLACK
    au.cell(i, 2, cell).font = BOLD
    au.cell(i, 3, prob).font = BLACK
    au.cell(i, 4, fix).font = BLACK
    for j in (2, 3, 4):
        au.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
    au.row_dimensions[i].height = 62

# ====================================================================== Lisez-moi
lm = wb.create_sheet("Lisez-moi", 0)
setw(lm, [110])
txt = [
    ("SEMELLE ISOLÉE SOUS POTEAU MÉTALLIQUE – RPA 2024 / BAEL 91 mod.99 / DTR BC 2.331", TITLE),
    ("", BLACK),
    ("Mode d'emploi", BOLD),
    ("1. « Données » : saisir les cellules bleues sur fond jaune (matériaux, sol, fût, A × B × h, zone / site / groupe).", BLACK),
    ("2. « Efforts » : coller les combinaisons (une ligne par combinaison et par appui), Type = ELU / ELS / ACC.", BLACK),
    ("   Option : « Combinaisons » génère ELU, ELS et les 8 combinaisons sismiques G + ψQ ± λ(Ex ± 0,3Ey) à partir des cas simples.", BLACK),
    ("3. « Vérifications » : portance, renversement, soulèvement, glissement pour chaque ligne (NON OK en rouge).", BLACK),
    ("4. « Ferraillage » : bielles et moments pour chaque ligne et chaque direction, poinçonnement.", BLACK),
    ("5. « Résultats » : taux maximaux, armatures retenues avec l'espacement imposé, ancrage, fût, dimensions conseillées.", BLACK),
    ("6. « Longrine » : effort RPA 2024 et ferraillage pour une portée libre.", BLACK),
    ("Ajuster A, B, h dans « Données » jusqu'à ce que tous les statuts de « Résultats » soient OK.", BOLD),
    ("", BLACK),
    ("Code couleur : bleu sur jaune = saisie ; noir = calcul ; vert / rouge = OK / NON OK.", BLACK),
    ("Unités : kN, kN·m, m, kPa (1 bar = 100 kPa), MPa, cm².", BLACK),
    ("Conventions : N > 0 compression ; A selon x (My), B selon y (Mx) ; moments ramenés à la sous-face : "
     "My,b = My + Vx·H, Mx,b = Mx − Vy·H, H = D + h0.", BLACK),
    ("", BLACK),
    ("Méthode identique à l'application « Semelles isolées – RPA 2024 » (vérification croisée possible). "
     "Différence : à l'ELS l'application résout exactement l'équilibre de la section ; ce classeur prend z = 0,9·d (sécuritaire).", GREY),
    ("Coefficients A, I, S, ψ : RPA 2024 tableaux 3.3, 3.11, 3.4/3.5 et 4.2.", GREY),
    ("L'onglet « Audit fichier d'origine » liste les erreurs trouvées dans l'ancien classeur.", GREY),
]
for i, (t, f) in enumerate(txt, 1):
    c = lm.cell(i, 1, t)
    c.font = f
    c.alignment = Alignment(wrap_text=True)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title not in ("Lisez-moi",)
    for row_ in ws.iter_rows():
        for c in row_:
            if c.font and c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)

# ====================================================================== Fiche semelle (style of the user's sheets)
fi = wb.create_sheet("Fiche semelle", 1)
TNR = "Times New Roman"
dbl = Side(style="double", color="000000")
DBL = Border(left=dbl, right=dbl, top=dbl, bottom=dbl)
F_GREY = PatternFill("solid", fgColor="BFBFBF")
F_YEL = PatternFill("solid", fgColor="FFFF00")
F_GRN = PatternFill("solid", fgColor="00B050")
F_LIGHT = PatternFill("solid", fgColor="F2F2F2")
DATA = {}  # fill colour → data cells (formatted only when their row holds a combination)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
fi.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJ", (3, 30, 22, 40, 18, 22, 40, 18, 16, 16)):
    fi.column_dimensions[col].width = w


def fc(addr, v, fill=None, fmt=None, color="000000", size=12, border=True, data=False):
    c = fi[addr]
    c.value = v
    c.font = Font(name=TNR, size=size, bold=True, color=color)
    c.alignment = CEN
    if data:
        DATA.setdefault(fill.fgColor.rgb if fill else None, []).append(addr)
    else:
        if border:
            c.border = DBL
        if fill:
            c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def fmerge(rng, v, fill=F_GREY, **kw):
    fi.merge_cells(rng)
    a = rng.split(":")[0]
    for row_ in fi[rng]:
        for c in row_:
            c.border = DBL
            if fill:
                c.fill = fill
    return fc(a, v, fill, **kw)


# --- header block (materials, soil, dimensions, weights)
for a, v in (("B1", "Fe"), ("C1", "FC28"), ("D1", "Ancrage (m)"), ("E1", "sigma sol (Bar)"), ("F1", "gamma sol (KN/m3)"),
             ("G1", "Poids terres (KN)"), ("G2", "Poids semelle (KN)"), ("G3", "Poids du fût (KN)"), ("G4", "NTOTAL (KN)"),
             ("D3", "A (m)"), ("E3", "B (m)"), ("F3", "H (m)"), ("B4", "a (m)"), ("B5", "b (m)")):
    fc(a, v, F_GREY)
fmerge("B3:C3", "Dimensions du poteau")
fc("B2", "=fe", fmt="0")
fc("C2", "=f_c28", fmt="0")
fc("D2", "=Dan", fmt="0.00")
fc("E2", "=sadm", fmt="0.00")
fc("F2", "=gr", fmt="0.00")
fc("H1", "=Ws", fmt="0.00")
fc("H2", "=Wf", fmt="0.00")
fc("H3", "=Wp", fmt="0.00")
fc("H4", "=Wtot", fmt="0.00")
fc("C4", "=aP", fmt="0.00")
fc("C5", "=bP", fmt="0.00")
fc("D4", "=Asem", F_YEL, "0.00")
fc("E4", "=Bsem", F_YEL, "0.00")
fc("F4", "=hsem", F_YEL, "0.00")
fc("D5", "=A_cons", fmt="0.00")
fc("E5", "=A_cons*bP/aP", fmt="0.00")
fc("F5", "=h_min", fmt="0.00")
fc("G5", "Ligne 5 : A, B conseillés et h minimal", F_LIGHT, size=9)
fc("H5", '=IF(AND(Asem>=D5-0.0001,hsem>=h_min-0.0001),"OK","NON OK")', color="FF0000")

# --- reactions and moments at the footing base
NF = 16  # first combinations of the Efforts sheet
r0 = 7
fc(f"B{r0}", "S1", F_GREY)
fmerge(f"C{r0}:E{r0}", "Réactions")
fmerge(f"F{r0}:J{r0}", "MOMENTS RAMENÉS À LA SOUS-FACE DE LA SEMELLE")
for j, h in enumerate(["Combinaison", "Fz (KN)", "Mx (KN.m)", "My (KN.m)", "Ntot (KN)", "My,b (KN.m)", "Mx,b (KN.m)",
                       "ex (m)", "ey (m)"], 2):
    fc(f"{get_column_letter(j)}{r0 + 1}", h, F_GREY)
for i in range(NF):
    r, e = r0 + 2 + i, 6 + i
    E = f"Efforts!$D{e}"
    fc(f"B{r}", f'=IF({E}="","",Efforts!B{e})', F_LIGHT, data=True)
    fc(f"C{r}", f'=IF({E}="","",{E})', fmt="0.00", data=True)
    fc(f"D{r}", f'=IF({E}="","",Efforts!G{e})', F_GRN, "0.00", data=True)
    fc(f"E{r}", f'=IF({E}="","",Efforts!H{e})', fmt="0.00", data=True)
    fc(f"F{r}", f"=Vérifications!E{e}", F_GRN, "0.00", data=True)
    fc(f"G{r}", f"=Vérifications!G{e}", fmt="0.00", data=True)
    fc(f"H{r}", f"=Vérifications!H{e}", fmt="0.00", data=True)
    fc(f"I{r}", f"=Vérifications!I{e}", fmt="0.000", data=True)
    fc(f"J{r}", f"=Vérifications!J{e}", fmt="0.000", data=True)

# --- soil stresses, X and Y
r1 = r0 + 2 + NF + 1
fmerge(f"C{r1}:F{r1}", 'Vérification des contraintes du sol "Sens_X"')
fmerge(f"G{r1}:I{r1}", 'Vérification des contraintes du sol "Sens_Y"')
for j, h in enumerate(["Combinaison", "Ntot (KN)", "Diagramme", "σ 3/4 (Bar)", "Vérification", "Diagramme", "σ 3/4 (Bar)",
                       "Vérification"], 2):
    fc(f"{get_column_letter(j)}{r1 + 1}", h, F_GREY)
for i in range(NF):
    r, e = r1 + 2 + i, 6 + i
    V = f"Vérifications!"
    fc(f"B{r}", f"=B{r0 + 2 + i}", F_LIGHT, data=True)
    fc(f"C{r}", f"={V}E{e}", F_GRN, "0.00", data=True)
    for cd, cs, col_d, col_s, L_ in (("D", "E", "K", "L", "A"), ("G", "H", "M", "N", "B")):
        fc(f"{cd}{r}", f'=IF({V}{col_d}{e}="trapèze","e < {L_}/6 => Diagramme Trapézoïdal",'
                       f'IF({V}{col_d}{e}="triangle","e > {L_}/6 => Diagramme triangulaire",{V}{col_d}{e}))', size=11, data=True)
        fc(f"{cs}{r}", f'=IF({V}{col_s}{e}="","",{V}{col_s}{e}/100)', fmt="0.00", data=True)
    fc(f"F{r}", f'=IF({V}L{e}="","",IF({V}L{e}<={V}O{e}+0.001,"OK","NON OK"))', color="FF0000", data=True)
    fc(f"I{r}", f'=IF({V}N{e}="","",IF({V}N{e}<={V}O{e}+0.001,"OK","NON OK"))', color="FF0000", data=True)

# --- stability
r2 = r1 + 2 + NF + 1
fmerge(f"C{r2}:I{r2}", "Vérification de la stabilité (renversement RPA 2024 §5.5 : Fs ≥ 1,3 sous séisme ; glissement §10.1.5)")
for j, h in enumerate(["Combinaison", "Taux portance", "Portance", "Taux renversement", "Renversement", "Taux glissement",
                       "Glissement", "Soulèvement"], 2):
    fc(f"{get_column_letter(j)}{r2 + 1}", h, F_GREY)
for i in range(NF):
    r, e = r2 + 2 + i, 6 + i
    fc(f"B{r}", f"=B{r0 + 2 + i}", F_LIGHT, data=True)
    for cl, src, f_ in (("C", "P", "0%"), ("D", "Q", None), ("E", "R", "0%"), ("F", "S", None), ("G", "X", "0%"),
                        ("H", "Y", None), ("I", "U", None)):
        fc(f"{cl}{r}", f"=Vérifications!{src}{e}", fmt=f_, color="000000" if f_ else "FF0000", data=True)

# --- reinforcement
r3 = r2 + 2 + NF + 1
fmerge(f"C{r3}:F{r3}", "Ferraillage de la semelle")
for j, h in enumerate(["", "As requis (cm²)", "Ferraillage retenu", "Espacement ≤ max", "Nappe sup. (cm²)"], 2):
    fc(f"{get_column_letter(j)}{r3 + 1}", h, F_GREY)
for k, (lab, col) in enumerate((("// A (x)", "B"), ("// B (y)", "C"))):
    r = r3 + 2 + k
    fc(f"B{r}", lab, F_GREY)
    fc(f"C{r}", f"=Résultats!{col}{r_as}", F_GRN, "0.00")
    fc(f"D{r}", f"=Résultats!{col}{r_txt}", F_YEL)
    fc(f"E{r}", f"=Résultats!{col}{r_txt + 1}", color="FF0000")
    fc(f"F{r}", f"=Résultats!{col}{r_top}", fmt="0.00")
from openpyxl.styles.differential import DifferentialStyle  # noqa: E402
for rgb, cells in DATA.items():
    for a in cells:
        row_ = "".join(ch for ch in a if ch.isdigit())
        kw = dict(border=DBL)
        if rgb:
            kw["fill"] = PatternFill("solid", fgColor=rgb[-6:], bgColor=rgb[-6:])
        fi.conditional_formatting.add(a, FormulaRule(formula=[f'LEN($B${row_})>0'], **kw))
for rg in (f"B9:J{r3 + 3}",):
    fi.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"NON OK"'], fill=PatternFill("solid", fgColor="FFC7CE")))
fi.freeze_panes = "A7"
fi.page_setup.orientation = "landscape"
fi.page_setup.fitToWidth = 1
fi.sheet_properties.pageSetUpPr.fitToPage = True

wb.active = 1
wb.save(OUT)
print("saved", OUT)
