"""Render a Trace as a Word note de calcul (python-docx)."""
from __future__ import annotations

from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .trace import Figure, Para, Step, Table, Trace

GREEN = RGBColor(0x00, 0x80, 0x3C)
RED = RGBColor(0xC0, 0x00, 0x00)
GREY = RGBColor(0x60, 0x60, 0x60)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
MATH = "Cambria Math"


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _run(p, text, bold=False, italic=False, size=None, color=None, font=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def _page_number(paragraph):
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = paragraph.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)


def _step(doc, s: Step):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    _run(p, "▪ " + s.label, bold=True, size=10)
    if s.ref:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(0.6)
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.keep_with_next = True
        _run(q, "Réf. : " + s.ref, italic=True, size=8, color=GREY)
    lines = []
    if s.sym:
        lines.append((s.sym, False))
    def pre(text, cond):
        return ("= " if cond and "=" not in text and "→" not in text[:3] else "") + text
    if s.num:
        lines.append((pre(s.num, s.sym), False))
    if s.res:
        lines.append((pre(s.res, s.sym or s.num), True))
    for text, bold in lines:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(1.0)
        q.paragraph_format.space_after = Pt(0)
        _run(q, text, bold=bold, size=10, font=MATH, color=BLUE if bold else None)
    if s.ok is not None:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(1.0)
        q.paragraph_format.space_after = Pt(2)
        if s.cond:
            _run(q, "Condition : " + s.cond + "  →  ", size=10, font=MATH)
        _run(q, "VÉRIFIÉE ✔" if s.ok else "NON VÉRIFIÉE ✘", bold=True, size=10, color=GREEN if s.ok else RED)


def _table(doc, t: Table):
    if t.title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        _run(p, t.title, bold=True, italic=True, size=9)
    tab = doc.add_table(rows=1, cols=len(t.headers))
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    small = len(t.headers) > 6 or len(t.rows) > 30
    fs = 7 if small else 9
    for i, h in enumerate(t.headers):
        c = tab.rows[0].cells[i]
        c.text = ""
        _run(c.paragraphs[0], str(h), bold=True, size=fs)
        _shade(c, "D9E2F3")
    for row in t.rows:
        cells = tab.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            txt = str(v)
            color = RED if txt.startswith("NON OK") else (GREEN if txt == "OK" else None)
            _run(cells[i].paragraphs[0], txt, size=fs, color=color, bold=color is not None)
    doc.add_paragraph()


def build_docx(tr: Trace, meta: dict) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    hdr = sec.header.paragraphs[0]
    _run(hdr, f"{meta.get('bureau', '')}   |   {meta.get('ouvrage', '')}   |   Note de calcul – Semelles isolées",
         size=8, color=GREY)
    ftr = sec.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(ftr, "Page ", size=8, color=GREY)
    _page_number(ftr)

    # ---- cover page
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "NOTE DE CALCUL", bold=True, size=26, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "Fondations superficielles – Semelles isolées sous poteaux métalliques", bold=True, size=14)
    doc.add_paragraph()
    tab = doc.add_table(rows=0, cols=2)
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in (("Projet / ouvrage", meta.get("ouvrage", "")), ("Maître d'ouvrage", meta.get("moa", "")),
                 ("Localisation", meta.get("lieu", "")), ("Bureau d'études", meta.get("bureau", "")),
                 ("Établi par", meta.get("ingenieur", "")), ("Date", meta.get("date", date.today().strftime("%d/%m/%Y"))),
                 ("Indice", meta.get("indice", "A")),
                 ("Règlements", "RPA 2024 – BAEL 91 mod.99 – CBA 93 – DTR BC 2.331 – CCM 97 – Eurocode 7")):
        c = tab.add_row().cells
        c[0].text, c[1].text = "", ""
        _run(c[0].paragraphs[0], k, bold=True)
        _shade(c[0], "D9E2F3")
        _run(c[1].paragraphs[0], str(v))
    doc.add_page_break()

    # ---- table of contents (Word updates it with F9)
    p = doc.add_paragraph()
    _run(p, "SOMMAIRE", bold=True, size=14, color=BLUE)
    p = doc.add_paragraph()
    r = p.add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        elif text:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            r._r.append(el)
        else:
            t = OxmlElement("w:t")
            t.text = "Clic droit → « Mettre à jour les champs » pour afficher le sommaire."
            r._r.append(t)
    doc.add_page_break()

    for s in tr.sections:
        if s.title:
            if s.level == 1 and s.title[:1].isdigit() and s.title.split(".")[0].isdigit() and int(s.title.split(".")[0]) >= 5:
                doc.add_page_break()
            h = doc.add_heading(s.title, level=min(s.level, 3))
            for r in h.runs:
                r.font.color.rgb = BLUE
        for it in s.items:
            if isinstance(it, Step):
                _step(doc, it)
            elif isinstance(it, Para):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                if it.kind == "warn":
                    _run(p, "⚠ " + it.text, bold=True, color=RED)
                elif it.kind == "note":
                    _run(p, "Note : " + it.text, italic=True, color=GREY, size=9)
                elif it.kind == "bold":
                    _run(p, it.text, bold=True)
                else:
                    _run(p, it.text)
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif isinstance(it, Table):
                _table(doc, it)
            elif isinstance(it, Figure):
                doc.add_picture(BytesIO(it.png), width=Cm(it.width_cm))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if it.caption:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _run(p, it.caption, italic=True, size=9)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
